from io import BytesIO

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.services.skills.builtin_tools import _fill_docx_template, _fill_xlsx_template


def test_fill_docx_template_preserves_template_and_fills_supported_patterns():
    source = Document()
    source.add_paragraph("项目：{{项目名称}}")
    row = source.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "负责人："
    row.cells[1].text = ""
    source_stream = BytesIO()
    source.save(source_stream)

    result, applied = _fill_docx_template(
        source_stream.getvalue(), {"项目名称": "智能填表", "负责人": "王小明"}
    )

    filled = Document(BytesIO(result))
    assert filled.paragraphs[0].text == "项目：智能填表"
    assert filled.tables[0].cell(0, 1).text == "王小明"
    assert set(applied) == {"项目名称", "负责人"}
    # The source bytes stay unchanged; the tool always returns a separate output.
    assert "{{项目名称}}" in Document(BytesIO(source_stream.getvalue())).paragraphs[0].text


def test_fill_xlsx_template_fills_placeholder_and_adjacent_blank_cell():
    source = Workbook()
    sheet = source.active
    sheet["A1"] = "项目：{{项目名称}}"
    sheet["A2"] = "负责人"
    source_stream = BytesIO()
    source.save(source_stream)

    result, applied = _fill_xlsx_template(
        source_stream.getvalue(), {"项目名称": "智能填表", "负责人": "王小明"}, {}
    )

    filled = load_workbook(BytesIO(result))
    sheet = filled.active
    assert sheet["A1"].value == "项目：智能填表"
    assert sheet["B2"].value == "王小明"
    assert set(applied) == {"项目名称", "负责人"}


def test_fill_xlsx_template_appends_rows_to_existing_header_table():
    source = Workbook()
    sheet = source.active
    sheet.title = "模型清单"
    sheet.append(["厂商", "模型名称", "发布时间"])
    sheet.append(["智谱", "GLM-4", "2025"])
    for cell in sheet[2]:
        cell.font = Font(bold=True)

    source_stream = BytesIO()
    source.save(source_stream)

    result, applied = _fill_xlsx_template(
        source_stream.getvalue(),
        {},
        {
            "模型清单": [
                {"厂商": "智谱", "模型名称": "GLM-5", "发布时间": "2026"},
                {"厂商": "月之暗面", "模型名称": "Kimi K2", "发布时间": "2025"},
            ]
        },
    )

    filled = load_workbook(BytesIO(result))
    sheet = filled["模型清单"]
    assert [sheet.cell(3, column).value for column in range(1, 4)] == ["智谱", "GLM-5", "2026"]
    assert [sheet.cell(4, column).value for column in range(1, 4)] == ["月之暗面", "Kimi K2", "2025"]
    assert sheet["A2"].value == "智谱"
    assert sheet["A3"].font.bold is True
    assert applied == ["模型清单（2行）"]


def test_fill_xlsx_template_inserts_before_summary_and_extends_simple_sum():
    source = Workbook()
    sheet = source.active
    sheet.title = "费用明细"
    sheet.append(["日期", "金额"])
    sheet.append(["2026-01-01", 10])
    sheet.append(["合计", "=SUM(B2:B2)"])
    source_stream = BytesIO()
    source.save(source_stream)

    result, applied = _fill_xlsx_template(
        source_stream.getvalue(),
        {},
        {"费用明细": [{"日期": "2026-01-02", "金额": 20}]},
    )

    filled = load_workbook(BytesIO(result))
    sheet = filled["费用明细"]
    assert [sheet.cell(3, column).value for column in range(1, 3)] == ["2026-01-02", 20]
    assert sheet["A4"].value == "合计"
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert applied == ["费用明细（1行）"]
