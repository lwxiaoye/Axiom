import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from PIL import Image

from app.schemas.schemas import ChatRequest
from app.services.files import document_parse_service, session_file_service, user_file_service


def _photo_bytes() -> bytes:
    image = Image.new("RGB", (180, 120))
    pixels = image.load()
    for y in range(image.height):
        for x in range(image.width):
            pixels[x, y] = ((x * 7 + y * 3) % 256, (x * 5) % 256, (y * 11) % 256)
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=94)
    return output.getvalue()


def test_chat_request_accepts_two_successful_images_with_null_notes():
    request = ChatRequest.model_validate({
        "message": "请对比这两张图",
        "attachments": [
            {"filename": "a.jpg", "kind": "image", "text": "A", "status": "ok", "note": None},
            {"filename": "b.jpg", "kind": "image", "text": "B", "status": "ok", "note": None},
        ],
    })

    assert request.attachments is not None
    assert [attachment.note for attachment in request.attachments] == ["", ""]


@pytest.mark.asyncio
async def test_parse_upload_skips_image_ocr_when_ocr_visual_disabled():
    called = {"ocr": False}

    async def boom(*_args, **_kwargs):
        called["ocr"] = True
        raise AssertionError("vision model uploads must not call OCR")

    with patch.object(document_parse_service, "_ocr_image", boom):
        parsed = await document_parse_service.parse_upload(
            "shot.png", _photo_bytes(), newapi_key="key", ocr_visual=False,
        )

    assert called["ocr"] is False
    assert parsed["kind"] == "image"
    assert parsed["status"] == "ok"
    assert parsed["text"] == ""
    assert "多模态" in str(parsed.get("note") or "")


@pytest.mark.asyncio
async def test_docx_embedded_photo_is_added_to_model_context():
    import docx

    document = docx.Document()
    document.add_paragraph("校园活动记录")
    document.add_picture(io.BytesIO(_photo_bytes()))
    output = io.BytesIO()
    document.save(output)

    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": True, "strategy": "multimodal_model", "model": "vision"}),
        ),
        patch.object(document_parse_service, "_describe_image", AsyncMock(return_value="照片中有学生在操场活动")),
    ):
        parsed = await document_parse_service.parse_upload("activity.docx", output.getvalue(), newapi_key="key")

    assert parsed["status"] == "ok"
    assert "校园活动记录" in parsed["text"]
    assert "照片中有学生在操场活动" in parsed["text"]


@pytest.mark.asyncio
async def test_pptx_embedded_photo_and_slide_text_are_both_parsed():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "班级活动"
    slide.shapes.add_picture(io.BytesIO(_photo_bytes()), Inches(1), Inches(1.5), width=Inches(3))
    output = io.BytesIO()
    presentation.save(output)

    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": True, "strategy": "multimodal_model", "model": "vision"}),
        ),
        patch.object(document_parse_service, "_describe_image", AsyncMock(return_value="照片展示了一场班级活动")),
    ):
        parsed = await document_parse_service.parse_upload("activity.pptx", output.getvalue(), newapi_key="key")

    assert parsed["kind"] == "pptx"
    assert parsed["status"] == "ok"
    assert "班级活动" in parsed["text"]
    assert "照片展示了一场班级活动" in parsed["text"]


def _distinct_blobs(n: int) -> list[tuple[str, bytes]]:
    """n 张互不相同、均大于 MIN_EMBEDDED_IMAGE_BYTES 的伪图片字节（视觉识别被 mock，无需真解码）。"""
    return [(f"img{i}.png", bytes([i]) + b"x" * 4000) for i in range(n)]


@pytest.mark.asyncio
async def test_embedded_images_over_cap_reports_pictures_not_slides():
    """超出上限时：note 用「张配图」措辞（不被误读成幻灯片页数）、正文说明文字已全读、状态 partial。"""
    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": True, "strategy": "multimodal_model", "model": "vision"}),
        ),
        patch.object(document_parse_service, "_describe_image", AsyncMock(return_value="配图描述")),
    ):
        body, status, note = await document_parse_service._ocr_embedded_images(
            _distinct_blobs(3), "key", "演示文稿", max_images=2
        )

    assert status == "partial"
    assert note == "仅识别前 2 张配图"
    assert "张配图" in note and "页" not in note  # 不出现「页」，杜绝「只读了前 N 页」的误读
    assert "正文文字已全部读取" in body


@pytest.mark.asyncio
async def test_embedded_images_within_cap_stay_ok():
    """未超上限时不降级：status=ok、note 为空、不出现配图告警。"""
    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": True, "strategy": "multimodal_model", "model": "vision"}),
        ),
        patch.object(document_parse_service, "_describe_image", AsyncMock(return_value="配图描述")),
    ):
        body, status, note = await document_parse_service._ocr_embedded_images(
            _distinct_blobs(3), "key", "演示文稿", max_images=25
        )

    assert status == "ok"
    assert note is None
    assert "配图较多" not in body


@pytest.mark.asyncio
async def test_large_document_retrieval_always_keeps_visual_description():
    text = "文档正文" * 1600 + "\n\n【文档中的图片内容（由视觉模型识别）】\n照片中是校园食堂"
    with patch.object(
        session_file_service.embedding_service,
        "get_active_embedding_config",
        AsyncMock(return_value=None),
    ):
        relevant = await session_file_service.retrieve_relevant("请总结文档", text)

    assert relevant.startswith("【文档中的图片内容")
    assert "照片中是校园食堂" in relevant


@pytest.mark.asyncio
async def test_subagent_file_read_can_see_stored_image_instead_of_binary_fallback():
    row = SimpleNamespace(id="file-1", filename="photo.png", mime="image/png", size_bytes=1024)
    parsed = {
        "kind": "image",
        "text": "【图片内容（由视觉模型识别）】\n校园操场",
        "truncated": False,
    }
    with (
        patch.object(user_file_service, "read_bytes", AsyncMock(return_value=(row, b"\x89PNG\x00binary"))),
        patch.object(document_parse_service, "parse_upload", AsyncMock(return_value=parsed)) as parse_mock,
    ):
        content = await user_file_service.get_content(
            "user-1", "file-1", newapi_key="key", ocr_embedded_images=True
        )

    assert content["kind"] == "image"
    assert "校园操场" in content["text"]
    parse_mock.assert_awaited_once()


def _scan_pdf_bytes() -> bytes:
    """纯图片 PDF（无文字层）= 扫描件。"""
    output = io.BytesIO()
    Image.new("RGB", (200, 120), "white").save(output, format="PDF")
    return output.getvalue()


@pytest.mark.asyncio
async def test_scanned_pdf_falls_back_to_vision_model_when_no_custom_endpoint():
    """线上只配了视觉模型、没有自建 OCR 端点：扫描件必须走视觉模型，而不是报「未配置端点」。"""
    seen: list[str] = []

    async def describe(content, ext, newapi_key, conf, *, audit_context=None):
        seen.append(ext)
        assert conf["strategy"] == "multimodal_model"
        return "重庆工程学院 图书馆 开放时间 8:00-22:00"

    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": True, "strategy": "multimodal_model", "model": "vision", "endpointUrl": ""}),
        ),
        patch.object(document_parse_service, "_describe_image", describe),
    ):
        parsed = await document_parse_service.parse_upload("scan.pdf", _scan_pdf_bytes(), newapi_key="key")

    assert seen == ["png"]
    assert parsed["status"] == "ok"
    assert parsed["text"].startswith("【扫描版 PDF（OCR 识别）】")
    assert "重庆工程学院" in parsed["text"] and "8:00" in parsed["text"]


@pytest.mark.asyncio
async def test_scanned_pdf_without_any_ocr_config_keeps_explicit_failure():
    """两条路都没配：仍是明确的 failed + 原因，不能悄悄去调一个没配模型名的视觉端点。"""
    with (
        patch.object(
            document_parse_service.cfg,
            "get_ocr_config",
            AsyncMock(return_value={"enabled": False, "strategy": "none", "model": "", "endpointUrl": ""}),
        ),
        patch.object(document_parse_service, "_describe_image", AsyncMock(side_effect=AssertionError("must not call vision"))),
    ):
        parsed = await document_parse_service.parse_upload("scan.pdf", _scan_pdf_bytes(), newapi_key="key")

    assert parsed["status"] == "failed"
    assert "未配置 OCR 端点" in str(parsed.get("note") or "")
