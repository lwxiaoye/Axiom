"""Harness: real-image PPT embed gate + capability broker freeze."""
from __future__ import annotations

import ast
import hashlib
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from app.services.chat.capability_broker import CapabilityBroker, build_capability_search_tool
from app.services.chat.tools.base import MainTool
from app.services.chat.tools.shell import _expected_images_from_attachments


SOURCE_FILE = (
    Path(__file__).resolve().parents[1] / "app/services/sandbox/output_review.py"
)


def _production_review_script() -> str:
    """与 test_ppt_output_review_policy 同一抽取口径，避免 placeholder 替换漂移。"""
    source = SOURCE_FILE.read_text(encoding="utf-8")
    tree = ast.parse(source)
    ranges = None
    points = None
    for node in tree.body:
        if not isinstance(node, ast.Assign) or len(node.targets) != 1:
            continue
        target = node.targets[0]
        if not isinstance(target, ast.Name):
            continue
        if target.id == "_PPT_FORBIDDEN_GLYPH_RANGES":
            ranges = ast.literal_eval(node.value)
        elif target.id == "_PPT_FORBIDDEN_GLYPH_POINTS":
            points = frozenset(ast.literal_eval(node.value.args[0]))
    match = re.search(r"_REVIEW_SCRIPT\s*=\s*r'''(.*?)'''\.replace", source, re.S)
    if ranges is None or points is None or match is None:
        raise AssertionError("无法从 output_review.py 提取 PPT 审查脚本")
    return (
        match.group(1)
        .replace("__PPT_FORBIDDEN_GLYPH_RANGES__", repr(ranges))
        .replace("__PPT_FORBIDDEN_GLYPH_POINTS__", repr(tuple(points)))
    )


def _tiny_png(color_byte: int = 0) -> bytes:
    base = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000a49444154789c6300000002000100d5fefcff0000000049454e44ae426082"
    )
    return base + bytes([color_byte, color_byte ^ 0xFF, 7, 3])


def _tool(name: str) -> MainTool:
    async def execute(_args):
        return "ok"

    return MainTool(name=name, description=f"{name} description", parameters={}, execute=execute)


class ExpectedImagesFromAttachmentsTests(unittest.TestCase):
    def test_extracts_image_attachments(self):
        names = _expected_images_from_attachments(
            [
                {"kind": "image", "filename": "photo.png"},
                {"kind": "text", "filename": "note.md"},
                {"filename": "shot.JPG"},
                {"kind": "image", "filename": "photo.png"},
            ]
        )
        self.assertEqual(names, ["photo.png", "shot.JPG"])


class BrokerFreezeTests(unittest.TestCase):
    def test_duplicate_pins_never_emit_duplicate_tool_schemas(self):
        tools = [_tool("bash"), _tool("fetch_ppt_asset"), _tool("publish_ppt_artifact")]
        broker = CapabilityBroker(
            tools,
            pinned=["bash", "fetch_ppt_asset", "fetch_ppt_asset", "publish_ppt_artifact"],
        )

        names = [tool.name for tool in broker.initial_tools()]
        self.assertEqual(names, ["bash", "fetch_ppt_asset", "publish_ppt_artifact"])

    def test_freeze_blocks_further_activate(self):
        tools = [_tool("ask_user_choice"), _tool("search_web"), _tool("browser_open"), _tool("bash")]
        broker = CapabilityBroker(tools, pinned=["ask_user_choice", "bash"])
        search = build_capability_search_tool(broker)
        broker.register(search, active=True)

        loaded = broker.activate("web search browser")
        self.assertTrue(loaded)
        active_after = set(broker.active_names)

        broker.freeze()
        self.assertTrue(broker.frozen)
        self.assertEqual(broker.activate("browser"), [])
        self.assertEqual(set(broker.active_names), active_after)
        text = broker.discovery_text("anything")
        self.assertIn("冻结", text)


class UserImagePptEmbedGateTests(unittest.TestCase):
    def test_review_script_rejects_placeholder_and_accepts_real_embed(self):
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            self.skipTest("当前测试环境未安装 python-pptx")

        script = _production_review_script()
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            user_png = root / "user-photo.png"
            user_bytes = _tiny_png(11)
            user_png.write_bytes(user_bytes)

            bad = Presentation()
            slide = bad.slides.add_slide(bad.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text = "假图交差"
            bad.save(root / "bad.pptx")

            good = Presentation()
            gslide = good.slides.add_slide(good.slide_layouts[6])
            gslide.shapes.add_picture(str(user_png), Inches(0.5), Inches(0.5), width=Inches(2))
            gbox = gslide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(4), Inches(1))
            gbox.text = "真实用户图"
            good.save(root / "good.pptx")

            other = root / "other.png"
            other.write_bytes(_tiny_png(99))
            wrong = Presentation()
            wslide = wrong.slides.add_slide(wrong.slide_layouts[6])
            wslide.shapes.add_picture(str(other), Inches(0.5), Inches(0.5), width=Inches(2))
            wrong.save(root / "wrong.pptx")

            patched = script.replace('OUTPUTS = "/workspace/outputs"', f"OUTPUTS = {str(root)!r}")
            script_path = root / "review.py"
            script_path.write_text(patched, encoding="utf-8")

            expected = json.dumps(["user-photo.png"], ensure_ascii=False)
            targets = json.dumps(["bad.pptx", "good.pptx", "wrong.pptx"], ensure_ascii=False)
            completed = subprocess.run(
                [sys.executable, "-I", str(script_path), "20", str(root), targets, expected],
                check=True,
                capture_output=True,
                text=True,
            )
            files = {item["name"]: item for item in json.loads(completed.stdout)["files"]}

            self.assertEqual(files["bad.pptx"]["status"], "failed")
            bad_checks = {c["name"]: c for c in files["bad.pptx"]["checks"]}
            self.assertEqual(bad_checks["user_images_embedded"]["status"], "failed")
            self.assertIn("没有任何嵌入图片", bad_checks["user_images_embedded"]["message"])

            self.assertEqual(files["good.pptx"]["status"], "passed")
            good_checks = {c["name"]: c for c in files["good.pptx"]["checks"]}
            self.assertEqual(good_checks["user_images_embedded"]["status"], "passed")

            self.assertEqual(files["wrong.pptx"]["status"], "failed")
            wrong_checks = {c["name"]: c for c in files["wrong.pptx"]["checks"]}
            self.assertEqual(wrong_checks["user_images_embedded"]["status"], "failed")
            self.assertIn("未以原始像素嵌入", wrong_checks["user_images_embedded"]["message"])

            import zipfile

            media_hashes = set()
            with zipfile.ZipFile(root / "good.pptx") as zf:
                for info in zf.infolist():
                    if info.filename.startswith("ppt/media/") and not info.is_dir():
                        media_hashes.add(hashlib.sha256(zf.read(info.filename)).hexdigest())
            self.assertIn(hashlib.sha256(user_bytes).hexdigest(), media_hashes)


if __name__ == "__main__":
    unittest.main()


def test_scrub_false_tool_outage_when_tools_delivered():
    from app.services.chat.turn_finalizer import (
        scrub_false_tool_outage_claim,
        tools_delivered_artifacts,
        tools_ran_successfully,
    )
    bad = (
        "当前轮次没有可用的命令执行入口，无法实际运行 bash 写入文件并回显。"
        "这条任务需要在下一轮重试。"
    )
    assert tools_delivered_artifacts([
        {"name": "bash", "status": "succeeded", "semantic_tags": ["artifact_producer"],
         "files": [{"id": "f1", "filename": "x.txt"}]}
    ])
    scrubbed = scrub_false_tool_outage_claim(bad, tools_succeeded=True)
    assert "命令执行入口" not in scrubbed
    assert "下一轮重试" not in scrubbed
    assert "已保存" in scrubbed or "已完成" in scrubbed
    assert scrub_false_tool_outage_claim(bad, tools_succeeded=False) == bad

    # write_file completed without top-level files must still count as delivered
    assert tools_delivered_artifacts([
        {"name": "write_file", "status": "completed", "semantic_tags": ["artifact_producer"],
         "preview": "已新建 a.md，已保存到「我的文件」。"}
    ])
    # observation.artifacts path
    assert tools_delivered_artifacts([
        {
            "name": "write_file",
            "status": "completed",
            "semantic_tags": ["artifact_producer"],
            "observation": {"status": "succeeded", "artifact_refs": [{"filename": "a.md"}]},
        }
    ])
    file_bad = (
        "当前没有文件读写操作入口，无法实际创建 live-m0-diag.md，"
        "不能谎称已保存。请在具备工具能力的下一轮重试。"
    )
    scrubbed2 = scrub_false_tool_outage_claim("收到。" + file_bad, tools_succeeded=True)
    assert "文件读写操作入口" not in scrubbed2
    assert "下一轮重试" not in scrubbed2

    # read tools success should scrub "无法访问文件区"
    assert tools_ran_successfully([{"name": "glob", "status": "completed", "preview": "ok"}])
    scrubbed3 = scrub_false_tool_outage_claim(
        "本轮我无法访问你的文件区，因此不能真正列出文件。请在新一轮对话中重试。",
        tools_succeeded=True,
    )
    assert "无法访问你的文件区" not in scrubbed3
    assert "重试" not in scrubbed3

def test_tools_success_blocks_plain_fallback_semantics():
    """工具已成功时不得再套 FALLBACK_NO_TOOLS 语义（与 main_tool_turn 守卫对齐）。"""
    from app.services.chat.types import TurnOutcome
    from app.services.chat.plain_turn import FALLBACK_NO_TOOLS_GUARD
    from app.services.chat.turn_finalizer import scrub_false_tool_outage_claim

    out = TurnOutcome()
    out.any_tool_succeeded = True
    out.write_tool_succeeded = True
    # 守卫条件：streamed_any or plan or tools_already
    tools_already = bool(out.get("any_tool_succeeded") or out.get("write_tool_succeeded"))
    assert tools_already is True
    assert out["streamed_any"] is False
    # 即便模型或回退护栏写出假故障，交付成功后仍 scrub
    bad = FALLBACK_NO_TOOLS_GUARD + "请在下一轮重试。"
    cleaned = scrub_false_tool_outage_claim(bad, tools_succeeded=True)
    assert "操作入口" not in cleaned
    assert "下一轮重试" not in cleaned



def test_scrub_ack_and_sandbox_outage_variants():
    from app.services.chat.turn_finalizer import scrub_false_tool_outage_claim

    assert scrub_false_tool_outage_claim(
        "收到。", tools_succeeded=True, tools_delivered=True
    ) == "已完成操作，文件已保存到「我的文件」。"
    assert scrub_false_tool_outage_claim(
        "当前沙箱未启用，请下一轮重试。", tools_succeeded=True, tools_delivered=True
    ) == "已完成操作，文件已保存到「我的文件」。"
    cleaned = scrub_false_tool_outage_claim(
        "文件已写好。作为文本模型我无法访问你的文件。",
        tools_succeeded=True,
        tools_delivered=True,
    )
    assert "文本模型" not in cleaned
    assert "文件已写好" in cleaned



def test_scrub_search_only_success_does_not_claim_file_saved():
    """search/glob 成功 ≠ 写产物：bare 确认不能 scrub 成「文件已保存」。"""
    from app.services.chat.turn_finalizer import scrub_false_tool_outage_claim

    out = scrub_false_tool_outage_claim(
        "收到。",
        tools_succeeded=True,
        tools_delivered=False,
    )
    assert "文件已保存" not in out
    assert out == "相关操作已完成。"

    # 假故障句也只 scrub 成非写产物口径
    out2 = scrub_false_tool_outage_claim(
        "当前没有可用的文件操作入口，请下一轮重试。",
        tools_succeeded=True,
        tools_delivered=False,
    )
    assert "文件已保存" not in out2
    assert out2 == "相关操作已完成。"


def test_strip_leading_mechanical_ack():
    from app.services.chat.turn_finalizer import strip_leading_mechanical_ack

    assert strip_leading_mechanical_ack("收到。8") == "8"
    assert strip_leading_mechanical_ack("好的，文件已写好。") == "文件已写好。"
    # 整段只是确认词：保留（用户可能要求单字）
    assert strip_leading_mechanical_ack("好") == "好"
    assert strip_leading_mechanical_ack("收到。") == "收到。"
    assert strip_leading_mechanical_ack("Python 特点如下") == "Python 特点如下"
    # v2.45：弱模型常见「我来帮你…」机械开场
    assert strip_leading_mechanical_ack("好的，我来帮你生成计算器页面。已写好。") == "已写好。"
    assert strip_leading_mechanical_ack("我来帮你查一下漳州天气。今天晴。") == "今天晴。"
    assert strip_leading_mechanical_ack("让我来处理一下。文件已保存。") == "文件已保存。"
    assert strip_leading_mechanical_ack("我先查漳州实时天气，同时搜东山岛的照片。漳州今天晴。") == "漳州今天晴。"
    assert strip_leading_mechanical_ack("当然，漳州今天晴。") == "漳州今天晴。"
    assert strip_leading_mechanical_ack("好嘞，已写好。") == "已写好。"
    assert strip_leading_mechanical_ack("我先从现有材料入手，准备检索。今天晴。") == "今天晴。"


def test_scrub_strips_ack_prefix_when_tools_ok():
    from app.services.chat.turn_finalizer import scrub_false_tool_outage_claim

    out = scrub_false_tool_outage_claim(
        "收到。已写入 live.txt",
        tools_succeeded=True,
        tools_delivered=True,
    )
    assert out.startswith("已写入")
    assert "收到" not in out


def test_scrub_contradictory_completion_v257():
    from app.services.chat.turn_finalizer import scrub_contradictory_completion, strip_leading_mechanical_ack

    raw = (
        "已创建 staged.md，第二节结论目前仍是 TODO，尚未填写实际内容。"
        "任务已全部完成，没有未执行的步骤。"
    )
    out = scrub_contradictory_completion(raw)
    assert "TODO" in out or "尚未" in out
    assert "任务已全部完成" not in out
    assert "没有未执行的步骤" not in out

    clean = "已写好 note.md，内容是 ok。"
    assert scrub_contradictory_completion(clean) == clean

    # strip expansions
    assert strip_leading_mechanical_ack("接下来查一下漳州天气。今天晴。") == "今天晴。"
    assert "今天晴" in strip_leading_mechanical_ack("稍等我先看一下。今天晴。")


def test_plan_humanize_and_toolish_v257():
    from app.services.agent_harness.model_driver import _humanize_toolish_plan_steps, _plan_titles_are_toolish, _provisional_plan_steps

    steps = _provisional_plan_steps(["search_web", "write_file"])
    titles = [s["title"] for s in steps]
    assert "查找需要的信息" in titles
    assert "保存到我的文件" in titles
    assert "收集资料" not in titles

    human = _humanize_toolish_plan_steps([
        {"title": "加载所需技能", "status": "pending"},
        {"title": "检索资料", "status": "pending"},
    ])
    assert human[0]["title"] == "准备要用的能力"
    assert human[1]["title"] == "查找需要的信息"
    assert _plan_titles_are_toolish([{"title": "加载所需技能"}])



def test_streaming_strip_holds_pure_opener_v257():
    from app.services.chat.turn_finalizer import StreamingMechanicalStripper, strip_leading_mechanical_ack
    assert strip_leading_mechanical_ack("我来查一下漳州天气，同时找几张东山岛的实景照片。") == ""
    st = StreamingMechanicalStripper()
    assert st.feed("我来查一下漳州天气，同时找几张东山岛的实景照片。") == ""
    out = st.feed("漳州今天晴，26℃。")
    assert out.startswith("漳州今天")
    assert "我来查" not in out


def test_strip_process_plan_residual_v301():
    """真机：机械开场被 stream 剥掉后，「先建立研究计划…」不得当正文。"""
    from app.services.chat.turn_finalizer import (
        StreamingMechanicalStripper,
        peel_commentary_from_answer,
        strip_leading_mechanical_ack,
    )

    full = "我来系统调研一下。先建立研究计划，然后多角度检索资料。"
    residual = "先建立研究计划，然后多角度检索资料。"
    assert strip_leading_mechanical_ack(full) == ""
    assert strip_leading_mechanical_ack(residual) == ""

    st = StreamingMechanicalStripper()
    assert st.feed(full) == ""
    assert st.flush() == ""

    # stream 已把残段当 delta 放出时，commentary 全文 endsWith 会失败——peel 必须清空
    assert peel_commentary_from_answer(residual, full) == ""
    assert peel_commentary_from_answer(full, full) == ""
    assert peel_commentary_from_answer("结论：人体工学椅。", full) == "结论：人体工学椅。"
    # 真答后缀不得被误剥
    assert peel_commentary_from_answer("前缀。" + residual, residual) == "前缀。"


def _openai_tool(name: str) -> dict:
    return {"type": "function", "function": {"name": name}}


def test_forced_product_tools_keep_investigate_when_plan_cursor_requires_it():
    from app.services.agent_harness.model_driver import _forced_product_payload_tools

    async def _execute(_args):
        return "ok"

    bash = MainTool(
        name="bash",
        description="bash",
        parameters={},
        execute=_execute,
        effect_scope="user_files",
        semantic_tags=("productive", "artifact_producer"),
    )
    publish = MainTool(
        name="publish_ppt_artifact",
        description="publish",
        parameters={},
        execute=_execute,
        capability="artifact.export",
        effect_scope="user_files",
    )
    search = MainTool(
        name="search_web",
        description="search",
        parameters={},
        execute=_execute,
        readonly=True,
        semantic_tags=("investigate", "web_search"),
    )
    fetch = MainTool(
        name="fetch_ppt_asset",
        description="fetch",
        parameters={},
        execute=_execute,
        effect_scope="scratch",
        semantic_tags=("download",),
    )
    tool_map = {tool.name: tool for tool in (bash, publish, search, fetch)}
    payload = [_openai_tool(name) for name in tool_map]
    profile = {
        "id": "artifact_coding",
        "artifact_kind": "presentation",
        "authoring_backend": "pptd",
        "qa_contract": {
            "image_requirement": {
                "mode": "searched_photos", "min_images": 3, "min_sources": 0, "brief": "",
            },
        },
    }
    plan_rows = [
        {"key": "step-1", "title": "理清演示结构", "status": "completed"},
        {
            "key": "step-2",
            "title": "搜索并下载库里比赛照片素材",
            "status": "in_progress",
            "requires": ["investigate"],
        },
    ]
    names = {
        ((item.get("function") or {}).get("name"))
        for item in _forced_product_payload_tools(
            payload, tool_map, plan_rows=plan_rows, execution_profile=profile, trace=[],
        )
    }
    assert "bash" in names
    assert "publish_ppt_artifact" in names
    assert "search_web" in names
    assert "fetch_ppt_asset" in names


def test_ppt_searched_photos_unmet_until_fetch_succeeds():
    from app.services.agent_harness.model_driver import _ppt_searched_photos_unmet

    profile = {
        "id": "artifact_coding",
        "artifact_kind": "presentation",
        "authoring_backend": "pptd",
        "qa_contract": {
            "image_requirement": {
                "mode": "searched_photos", "min_images": 3, "min_sources": 0, "brief": "",
            },
        },
    }
    assert _ppt_searched_photos_unmet(profile, []) is True
    assert _ppt_searched_photos_unmet(profile, [
        {"name": "fetch_ppt_asset", "status": "failed", "failed": True},
    ]) is True
    assert _ppt_searched_photos_unmet(profile, [
        {"name": "fetch_ppt_asset", "status": "succeeded"},
        {"name": "fetch_ppt_asset", "status": "succeeded"},
        {"name": "fetch_ppt_asset", "status": "succeeded"},
    ]) is False
