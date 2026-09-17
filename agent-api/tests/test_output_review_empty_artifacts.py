import contextlib
import io
import json
import sys

import openpyxl
from docx import Document
from pypdf import PdfWriter
from pptx import Presentation

from app.services.sandbox.output_review import _REVIEW_SCRIPT


def _review(tmp_path) -> dict:  # noqa: ANN001
    script = _REVIEW_SCRIPT.replace('OUTPUTS = "/workspace/outputs"', f"OUTPUTS = {str(tmp_path)!r}")
    stdout = io.StringIO()
    original_argv = sys.argv
    try:
        sys.argv = ["output_review", "20"]
        with contextlib.redirect_stdout(stdout):
            exec(script, {})  # noqa: S102 - runs the fixed in-module review script in a temp directory
    finally:
        sys.argv = original_argv
    return json.loads(stdout.getvalue())


def test_review_rejects_structurally_valid_but_empty_artifacts(tmp_path):
    Document().save(tmp_path / "empty.docx")
    Presentation().save(tmp_path / "empty.pptx")
    openpyxl.Workbook().save(tmp_path / "empty.xlsx")
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with (tmp_path / "empty.pdf").open("wb") as output:
        writer.write(output)
    (tmp_path / "empty.txt").write_text("  \n", encoding="utf-8")

    report = _review(tmp_path)
    statuses = {item["name"]: item["status"] for item in report["files"]}
    assert statuses == {
        "empty.docx": "failed",
        "empty.pdf": "failed",
        "empty.pptx": "failed",
        "empty.txt": "failed",
        "empty.xlsx": "failed",
    }
