from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.routers import workflow
from app.services.chat import builtin_app_access
from app.services.workflows import agent_conversation_log_service
from app.services.workflows.agent_conversation_log_service import build_conversation_log_xlsx, normalize_conversation_log_filters


def test_run_message_append_exposes_a_validated_turn_contract():
    payload = workflow.RunMessageAppend(
        sessionId="thread-1",
        role="assistant",
        content="reply",
        turnId="turn-1",
        status="failed",
    )

    assert payload.turnId == "turn-1"
    assert payload.status == "failed"

    with pytest.raises(ValidationError):
        workflow.RunMessageAppend(
            sessionId="thread-1",
            role="assistant",
            content="reply",
            status="anything",
        )


def test_run_message_reader_returns_persisted_turn_status_and_feedback():
    source = Path(workflow.__file__).read_text(encoding="utf-8")
    handler = source[source.index("async def run_messages"):source.index("async def append_run_message")]

    assert '"turnId": r.turn_id' in handler
    assert '"status": r.status' in handler
    assert '"feedback": r.feedback' in handler


def test_conversation_log_routes_are_available_for_admin_and_owner():
    source = Path(workflow.__file__).read_text(encoding="utf-8")

    assert '@router.get("/admin/app/{app_id}/conversation-logs")' in source
    assert '@router.get("/admin/app/{app_id}/conversation-logs/export")' in source
    assert '@router.get("/app/{app_id}/conversation-logs")' in source
    assert '@router.get("/app/{app_id}/conversation-logs/export")' in source


def test_conversation_log_service_has_safe_formula_escaping():
    service_path = Path(workflow.__file__).parents[1] / "services" / "workflows" / "agent_conversation_log_service.py"
    source = service_path.read_text(encoding="utf-8") if service_path.exists() else ""

    assert "def excel_text" in source
    assert '"=", "+", "-", "@"' in source


def test_conversation_log_filters_use_the_same_monitoring_range_keys():
    filters = normalize_conversation_log_filters(range_key="quarter_to_date")

    assert filters.range_key == "quarter_to_date"


def test_conversation_log_export_groups_session_columns_and_outputs_each_question_answer():
    content = build_conversation_log_xlsx("报修助手", [{
        "title": "空调报修",
        "username": "张三",
        "status": "completed",
        "messageCount": 4,
        "lastMessageAt": "2026-09-08T10:04:00",
        "upvotes": 1,
        "downvotes": 1,
        "messages": [
            {"role": "user", "content": "空调不制冷", "createdAt": "2026-09-08T10:00:00"},
            {"role": "assistant", "content": "已为您登记报修", "createdAt": "2026-09-08T10:01:00", "status": "completed", "feedback": "up"},
            {"role": "user", "content": "预计多久处理？", "createdAt": "2026-09-08T10:03:00"},
            {"role": "assistant", "content": "预计两小时内处理", "createdAt": "2026-09-08T10:04:00", "status": "completed", "feedback": "down"},
        ],
    }])

    sheet = load_workbook(BytesIO(content)).active

    assert [cell.value for cell in sheet[1]] == [
        "应用名称", "会话标题", "用户姓名", "会话状态", "消息数", "最近消息时间", "点赞数", "点踩数",
        "提问时间", "用户问题", "回答时间", "智能体回答", "回答状态", "用户反馈",
    ]
    assert sheet.max_row == 3
    assert sheet["J2"].value == "空调不制冷"
    assert sheet["L2"].value == "已为您登记报修"
    assert sheet["J3"].value == "预计多久处理？"
    assert sheet["L3"].value == "预计两小时内处理"
    assert sheet["F2"].value == "2026-09-08 10:04:00"
    assert sheet["I2"].value == "2026-09-08 10:00:00"
    assert sheet["K3"].value == "2026-09-08 10:04:00"
    assert "A2:A3" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert "H2:H3" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}


def test_conversation_log_time_uses_the_shared_page_and_export_format():
    assert agent_conversation_log_service.format_conversation_log_time("2026-09-08T11:52:20.541000") == "2026-09-08 11:52:20"


@pytest.mark.asyncio
async def test_message_feedback_access_allows_the_owner_run_session(monkeypatch):
    """运行页消息也在 ai_chat_messages，但不属于主对话/内置应用会话。"""
    thread = SimpleNamespace(id="run-thread", user_id="user-1", app_id="app-1")

    class Result:
        def scalar_one_or_none(self):
            return thread

    class Session:
        async def execute(self, _statement):
            return Result()

    class SessionFactory:
        def __call__(self):
            return self

        async def __aenter__(self):
            return Session()

        async def __aexit__(self, *_args):
            return None

    async def ordinary_thread_access(*_args, **_kwargs):
        raise AssertionError("运行会话不应走只允许主对话/内置应用的权限校验")

    monkeypatch.setattr(builtin_app_access, "async_session", SessionFactory())
    monkeypatch.setattr(builtin_app_access, "require_thread_access", ordinary_thread_access)

    actual = await builtin_app_access.require_message_access(SimpleNamespace(user_id="user-1"), 8850)

    assert actual is thread
