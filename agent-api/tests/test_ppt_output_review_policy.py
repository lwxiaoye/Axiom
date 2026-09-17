import ast
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


SOURCE_FILE = (
    Path(__file__).resolve().parents[1]
    / "app/services/sandbox/output_review.py"
)


def _production_policy() -> tuple[tuple, frozenset, str]:
    source = SOURCE_FILE.read_text(encoding="utf-8")
    tree = ast.parse(source)
    ranges = None
    points = None
    for node in tree.body:
        if not isinstance(node, ast.Assign) or len(node.targets) != 1:
            continue
        target = node.targets[0]
        if not isinstance(target, ast.Name):
            continue
        if target.id == "_PPT_FORBIDDEN_GLYPH_RANGES":
            ranges = ast.literal_eval(node.value)
        elif target.id == "_PPT_FORBIDDEN_GLYPH_POINTS":
            points = frozenset(ast.literal_eval(node.value.args[0]))
    match = re.search(
        r"_REVIEW_SCRIPT\s*=\s*r'''(.*?)'''\.replace",
        source,
        re.S,
    )
    if ranges is None or points is None or match is None:
        raise AssertionError("无法从 output_review.py 提取 PPT 兼容策略")
    script = match.group(1).replace(
        "__PPT_FORBIDDEN_GLYPH_RANGES__", repr(ranges)
    ).replace("__PPT_FORBIDDEN_GLYPH_POINTS__", repr(tuple(points)))
    return ranges, points, script


def _is_forbidden(char: str) -> bool:
    ranges, points, _ = _production_policy()
    codepoint = ord(char)
    return codepoint in points or any(
        start <= codepoint <= end for start, end in ranges
    )


class PptOutputReviewPolicyTests(unittest.TestCase):
    def test_forbidden_glyph_ranges_cover_product_regression(self) -> None:
        for char in "💧⚠✅⏰✓":
            self.assertTrue(_is_forbidden(char), char)
        for char in "补水ABC123•—%":
            self.assertFalse(_is_forbidden(char), char)

    def test_review_script_rejects_emoji_and_icon_fonts(self) -> None:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.util import Inches
        except ImportError:
            self.skipTest("当前测试环境未安装 python-pptx")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(6), Inches(1)
            )
            run = box.text_frame.paragraphs[0].add_run()
            run.text = "💧 多喝水"
            run.font.name = "Wingdings"
            prs.save(root / "bad.pptx")

            safe = Presentation()
            slide = safe.slides.add_slide(safe.slide_layouts[6])
            slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(1), Inches(1)
            )
            box = slide.shapes.add_textbox(
                Inches(2.2), Inches(1), Inches(6), Inches(1)
            )
            box.text = "多喝水，更健康"
            safe.save(root / "safe.pptx")

            _, _, production_script = _production_policy()
            script = production_script.replace(
                'OUTPUTS = "/workspace/outputs"', f"OUTPUTS = {str(root)!r}"
            )
            script_path = root / "review.py"
            script_path.write_text(script, encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, "-I", str(script_path), "20"],
                check=True,
                capture_output=True,
                text=True,
            )
            files = {
                item["name"]: item
                for item in json.loads(completed.stdout)["files"]
            }

            self.assertEqual(files["bad.pptx"]["status"], "failed")
            bad_checks = {item["name"]: item for item in files["bad.pptx"]["checks"]}
            self.assertEqual(bad_checks["web_safe_icons"]["status"], "failed")
            self.assertEqual(bad_checks["no_icon_fonts"]["status"], "failed")
            self.assertEqual(files["safe.pptx"]["status"], "passed")


if __name__ == "__main__":
    unittest.main()
