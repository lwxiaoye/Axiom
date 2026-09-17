"""文档排版 / 格式转换系统工具（builtin.document_typeset、builtin.document_convert，
以及 builtin.document_export 的模板 + PDF 扩展）。沙箱 LibreOffice 转换用桩替代。"""
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.services.files import user_file_service
from app.services.skills import builtin_tools, document_typeset
from app.services.skills.document_typeset import (
    TEMPLATES,
    TEMPLATE_FORMAL_REPORT,
    TEMPLATE_WORK_SUMMARY,
    convert_document,
    docx_to_markdown,
    normalize_heading_levels,
    outline_of,
    parse_blocks,
    pop_title,
    render_docx,
    resolve_template,
    typeset_document,
)

SAMPLE_MD = """# 计算机学院2025年度工作总结

2025年，学院围绕立德树人根本任务扎实推进各项工作。

## 工作回顾

### 教学工作
新增省级一流课程 **2 门**。

| 指标 | 2024年 | 2025年 |
| --- | --- | --- |
| 毕业率 | 96.2% | 97.5% |

## 下阶段安排
- 推进一流专业建设
1. 强化产教融合
"""

PLAIN_CN = """一、工作回顾
今年完成了教学改革任务。
（一）教学工作
新增课程两门。
1. 课程建设
建设了一批课程。
二、下阶段安排
继续推进。
"""


# ---------------------------------------------------------------- 解析


def test_parse_blocks_markdown_headings_tables_and_lists():
    blocks = parse_blocks(SAMPLE_MD)
    title = pop_title(blocks)
    normalize_heading_levels(blocks)

    assert title == "计算机学院2025年度工作总结"
    assert [item["level"] for item in outline_of(blocks)] == [1, 2, 1]
    assert [item["text"] for item in outline_of(blocks)] == ["工作回顾", "教学工作", "下阶段安排"]
    table = next(block for block in blocks if block.kind == "table")
    assert table.headers == ["指标", "2024年", "2025年"]
    assert table.rows == [["毕业率", "96.2%", "97.5%"]]
    lists = [block for block in blocks if block.kind == "list"]
    assert [(item.ordered, item.text) for item in lists] == [(False, "推进一流专业建设"), (True, "强化产教融合")]


def test_parse_blocks_falls_back_to_chinese_numbering_without_markdown_headings():
    blocks = parse_blocks(PLAIN_CN)
    assert [(item["level"], item["text"]) for item in outline_of(blocks)] == [
        (1, "一、工作回顾"),
        (2, "（一）教学工作"),
        (3, "1. 课程建设"),
        (1, "二、下阶段安排"),
    ]
    # 正文行不会被误判成标题
    assert any(block.kind == "paragraph" and block.text == "今年完成了教学改革任务。" for block in blocks)


def test_pop_title_keeps_first_heading_when_multiple_level1_exist():
    blocks = parse_blocks("# 一、背景\n正文\n# 二、方法\n正文")
    assert pop_title(blocks) is None
    assert len([block for block in blocks if block.kind == "heading"]) == 2


def test_resolve_template_accepts_aliases_and_rejects_unknown():
    assert resolve_template("工作总结").key == TEMPLATE_WORK_SUMMARY
    assert resolve_template("「正式报告」").key == TEMPLATE_FORMAL_REPORT
    assert resolve_template("formal-report").key == TEMPLATE_FORMAL_REPORT
    assert resolve_template("科研报告").key == TEMPLATE_FORMAL_REPORT
    assert resolve_template("红头文件") is None
    assert resolve_template("") is None


# ---------------------------------------------------------------- 渲染


def _render(template_key: str, organization="计算机学院", date="2026年1月10日") -> Document:
    blocks = parse_blocks(SAMPLE_MD)
    title = pop_title(blocks)
    normalize_heading_levels(blocks)
    data = render_docx(blocks, TEMPLATES[template_key], title=title, organization=organization, date=date)
    return Document(BytesIO(data))


def _east_asia_font(run) -> str:
    rfonts = run._element.rPr.find(qn("w:rFonts"))
    return rfonts.get(qn("w:eastAsia"))


def test_render_work_summary_applies_heading_styles_numbering_indent_and_signature():
    document = _render(TEMPLATE_WORK_SUMMARY)
    headings = [p for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert [(p.style.name, p.text) for p in headings] == [
        ("Heading 1", "一、工作回顾"),
        ("Heading 2", "（一）教学工作"),
        ("Heading 1", "二、下阶段安排"),
    ]
    assert _east_asia_font(headings[0].runs[0]) == "黑体"
    assert _east_asia_font(headings[1].runs[0]) == "楷体_GB2312"
    body = next(p for p in document.paragraphs if p.text.startswith("2025年"))
    assert _east_asia_font(body.runs[0]) == "仿宋_GB2312"
    assert body._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars")) == "200"
    assert body.runs[0].font.size.pt == 16
    # 表格：Table Grid + 表头加粗
    assert len(document.tables) == 1
    assert document.tables[0].style.name == "Table Grid"
    assert document.tables[0].cell(0, 0).paragraphs[0].runs[0].font.bold is True
    # 落款靠右
    tail = [p.text for p in document.paragraphs[-2:]]
    assert tail == ["计算机学院", "2026年1月10日"]
    # 页码域
    footer_xml = document.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "— " in document.sections[0].footer.paragraphs[0].text
    # 页边距按公文
    section = document.sections[0]
    assert round(section.top_margin.cm, 2) == 3.7 and round(section.left_margin.cm, 2) == 2.8


def test_render_formal_report_uses_decimal_numbering_and_meta_under_title():
    document = _render(TEMPLATE_FORMAL_REPORT)
    texts = [p.text for p in document.paragraphs]
    assert texts[:3] == ["计算机学院2025年度工作总结", "计算机学院", "2026年1月10日"]
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["1  工作回顾", "1.1  教学工作", "2  下阶段安排"]
    body = next(p for p in document.paragraphs if p.text.startswith("2025年"))
    assert _east_asia_font(body.runs[0]) == "宋体"
    assert body.runs[0].font.size.pt == 12
    # 正式报告不加文末落款
    assert not texts[-1].startswith("2026年")


def test_render_does_not_double_number_headings_that_already_carry_numbers():
    blocks = parse_blocks("## 一、背景\n正文\n## 二、方法\n正文")
    normalize_heading_levels(blocks)
    document = Document(BytesIO(render_docx(blocks, TEMPLATES[TEMPLATE_WORK_SUMMARY], title="t")))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["一、背景", "二、方法"]


def test_docx_to_markdown_roundtrips_headings_and_tables():
    document = _render(TEMPLATE_FORMAL_REPORT)
    stream = BytesIO()
    document.save(stream)
    markdown = docx_to_markdown(stream.getvalue())
    assert "# 1  工作回顾" in markdown
    assert "## 1.1  教学工作" in markdown
    assert "| 指标 | 2024年 | 2025年 |" in markdown
    assert "| --- | --- | --- |" in markdown


# ---------------------------------------------------------------- 回执


def test_pop_generated_file_receipts_merges_single_and_list_and_dedupes():
    result = {
        "ok": True,
        "_generated_file_receipt": {"id": "a"},
        "_generated_file_receipts": [{"id": "a"}, {"id": "b"}, {"id": ""}, "junk"],
    }
    receipts = builtin_tools.pop_generated_file_receipts(result)
    assert [item["id"] for item in receipts] == ["a", "b"]
    assert "_generated_file_receipt" not in result and "_generated_file_receipts" not in result
    assert builtin_tools.pop_generated_file_receipts("text") == []


# ---------------------------------------------------------------- 工具执行（沙箱与文件服务用桩）


class _FakeRow:
    def __init__(self, filename: str):
        self.filename = filename


@pytest.fixture
def fake_files(monkeypatch):
    saved: list[dict] = []
    store: dict[str, tuple[str, bytes]] = {}

    async def save_generated_bytes(user_id, filename, data, *, thread_id=None, run_id=None):
        saved.append({"filename": filename, "size": len(data), "data": data})
        return {"id": f"file-{len(saved)}", "filename": filename, "size": len(data), "versionNo": 1}

    async def read_bytes(user_id, file_id):
        if file_id not in store:
            raise ValueError("missing")
        name, data = store[file_id]
        return _FakeRow(name), data

    async def docx_to_pdf(filename, data):
        return b"%PDF-1.4 fake " + filename.encode("utf-8")

    async def soffice_convert(filename, data, target_ext):
        return f"converted:{target_ext}:{filename}".encode("utf-8")

    monkeypatch.setattr(user_file_service, "save_generated_bytes", save_generated_bytes)
    monkeypatch.setattr(user_file_service, "read_bytes", read_bytes)
    monkeypatch.setattr(document_typeset, "docx_to_pdf", docx_to_pdf)
    monkeypatch.setattr(document_typeset, "soffice_convert", soffice_convert)
    return {"saved": saved, "store": store}


RUNTIME = {"user_id": "u1", "thread_id": "t1", "run_id": "r1"}


@pytest.mark.asyncio
async def test_typeset_document_outputs_docx_and_pdf_with_outline(fake_files):
    result = await typeset_document(
        {"template": "工作总结", "content": SAMPLE_MD, "output_format": "docx+pdf", "organization": "计算机学院"},
        RUNTIME,
    )
    assert [item["filename"] for item in fake_files["saved"]] == [
        "计算机学院2025年度工作总结.docx",
        "计算机学院2025年度工作总结.pdf",
    ]
    assert fake_files["saved"][0]["data"][:2] == b"PK"
    assert result["template"] == "工作总结"
    assert result["heading_counts"] == {"level1": 2, "level2": 1, "level3": 0}
    assert result["table_count"] == 1
    assert [item["id"] for item in result["_generated_file_receipts"]] == ["file-1", "file-2"]
    assert result["_generated_file_receipt"]["id"] == "file-1"
    assert result["files"][1]["mime"] == "application/pdf"
    assert result["file"]["origin"]["tool"] == "builtin.document_typeset"


@pytest.mark.asyncio
async def test_typeset_document_reads_docx_source_from_my_files(fake_files):
    source = Document()
    source.add_heading("调研报告", level=1)
    source.add_heading("背景", level=2)
    source.add_paragraph("正文内容。")
    stream = BytesIO()
    source.save(stream)
    fake_files["store"]["src-1"] = ("调研.docx", stream.getvalue())

    result = await typeset_document({"template": "正式报告", "source_file_id": "src-1"}, RUNTIME)
    assert result["title"] == "调研报告"
    assert result["outline"] == [{"level": 1, "text": "背景"}]
    assert fake_files["saved"][0]["filename"] == "调研报告.docx"
    assert result["file"]["origin"]["sourceFileId"] == "src-1"


@pytest.mark.asyncio
async def test_typeset_document_rejects_unknown_template_and_empty_body(fake_files):
    with pytest.raises(ValueError, match="内置模板"):
        await typeset_document({"template": "红头文件", "content": "正文"}, RUNTIME)
    with pytest.raises(ValueError, match="正文为空"):
        await typeset_document({"template": "工作总结", "content": "  "}, RUNTIME)
    with pytest.raises(ValueError, match="只有标题"):
        await typeset_document({"template": "工作总结", "content": "# 只有标题"}, RUNTIME)
    assert fake_files["saved"] == []


@pytest.mark.asyncio
async def test_convert_document_docx_to_pdf_keeps_original_layout(fake_files):
    fake_files["store"]["doc-1"] = ("年度总结.docx", b"PK-docx-bytes")
    result = await convert_document({"source_file_id": "doc-1", "target_format": "pdf"}, RUNTIME)
    assert fake_files["saved"][0]["filename"] == "年度总结.pdf"
    assert fake_files["saved"][0]["data"] == "converted:pdf:年度总结.docx".encode("utf-8")
    assert result["source_format"] == "docx" and result["target_format"] == "pdf"
    assert result["_generated_file_receipt"]["origin"]["sourceFileId"] == "doc-1"


@pytest.mark.asyncio
async def test_convert_document_docx_to_markdown_and_markdown_to_docx(fake_files):
    source = Document()
    source.add_heading("标题", level=1)
    source.add_paragraph("段落。")
    stream = BytesIO()
    source.save(stream)
    fake_files["store"]["doc-2"] = ("说明.docx", stream.getvalue())
    result = await convert_document({"source_file_id": "doc-2", "target_format": "md"}, RUNTIME)
    assert fake_files["saved"][-1]["filename"] == "说明.md"
    assert fake_files["saved"][-1]["data"].decode("utf-8").startswith("# 标题")
    assert result["target_format"] == "md"

    fake_files["store"]["md-1"] = ("提纲.md", "# 提纲\n\n## 一、背景\n正文".encode("utf-8"))
    await convert_document({"source_file_id": "md-1", "target_format": "docx"}, RUNTIME)
    assert fake_files["saved"][-1]["filename"] == "提纲.docx"
    assert fake_files["saved"][-1]["data"][:2] == b"PK"

    with pytest.raises(ValueError, match="无需转换"):
        await convert_document({"source_file_id": "md-1", "target_format": "md"}, RUNTIME)
    fake_files["store"]["ppt-1"] = ("汇报.pptx", b"PK")
    with pytest.raises(ValueError, match="暂不支持"):
        await convert_document({"source_file_id": "ppt-1", "target_format": "pdf"}, RUNTIME)


@pytest.mark.asyncio
async def test_document_export_supports_template_and_pdf(fake_files):
    result = await builtin_tools.execute_builtin_tool(
        "builtin.document_export",
        {"filename": "总结", "format": "pdf", "template": "工作总结", "content": SAMPLE_MD},
        RUNTIME,
    )
    assert fake_files["saved"][0]["filename"] == "总结.pdf"
    assert fake_files["saved"][0]["data"].startswith(b"%PDF")
    assert result["file"]["mime"] == "application/pdf"
    assert result["file"]["origin"]["template"] == "工作总结"

    plain = await builtin_tools.execute_builtin_tool(
        "builtin.document_export",
        {"filename": "普通", "format": "docx", "content": "# 标题\n正文"},
        RUNTIME,
    )
    assert fake_files["saved"][1]["filename"] == "普通.docx"
    assert "template" not in plain["file"]["origin"]

    with pytest.raises(ValueError, match="template"):
        await builtin_tools.execute_builtin_tool(
            "builtin.document_export",
            {"filename": "x", "format": "docx", "template": "红头", "content": "正文"},
            RUNTIME,
        )


def test_builtin_tool_catalog_declares_new_tools():
    ids = {tool["id"] for tool in builtin_tools.BUILTIN_TOOLS}
    assert {"builtin.document_typeset", "builtin.document_convert", "builtin.document_export"} <= ids
    export = builtin_tools.BUILTIN_TOOL_MAP["builtin.document_export"]
    assert "pdf" in export["parameters"]["properties"]["format"]["enum"]
    assert export["parameters"]["properties"]["template"]["enum"] == ["默认", "工作总结", "正式报告"]
    typeset = builtin_tools.BUILTIN_TOOL_MAP["builtin.document_typeset"]
    assert typeset["parameters"]["properties"]["template"]["enum"] == ["工作总结", "正式报告"]
    assert typeset["parameters"]["required"] == ["template"]
