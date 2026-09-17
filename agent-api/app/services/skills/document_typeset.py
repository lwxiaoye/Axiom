"""
文档排版与格式转换（系统工具 builtin.document_typeset / builtin.document_convert 的实现，
并为 builtin.document_export 提供「套模板 + PDF」能力）。

招标口径（写作助手 / 文档排版）落到这里的三件事：
- 场景模板选择：平台内置「工作总结」「正式报告」两套实际可用模板（`TEMPLATES`），
  不再要求用户每次自己上传 DOCX 模板；用户自带占位符模板仍走 builtin.template_fill。
- 模板排版与格式调整：Markdown/纯文本正文 → 标题层级、正文字体字号、段落间距与首行缩进、
  基础表格、页码，全部由 `render_docx` 用 python-docx 真实写进 DOCX；预览复用「我的文件」
  既有的 PDF 预览链路。
- 多格式导入与导出：DOCX/DOC/MD/TXT 输入，DOCX/PDF/MD 输出。PDF 走沙箱 LibreOffice
  （与文件预览同一条 soffice 转换链路，沙箱镜像已含 Noto CJK 字体）。

字体按公文/报告惯例写入 w:eastAsia（仿宋_GB2312 / 黑体 / 楷体_GB2312 / 宋体）；用户机器缺字体时
Word 自动替换，沙箱转 PDF 时由 Noto CJK 兜底，层级、字号、缩进与行距不受影响。
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Optional

TEMPLATE_WORK_SUMMARY = "工作总结"
TEMPLATE_FORMAL_REPORT = "正式报告"
TEMPLATE_DEFAULT = "默认"

_SOFFICE_TIMEOUT_MS = 120_000

_MIMES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "md": "text/markdown",
}


# ---------------------------------------------------------------------------
# 模板规格
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _FontSpec:
    east: str          # 中文字体（w:eastAsia）
    ascii: str         # 西文/数字字体
    size: float        # 磅
    bold: bool = False


@dataclass(frozen=True)
class TemplateSpec:
    key: str
    summary: str                      # 给模型/用户看的一句话规格
    margins_cm: tuple[float, float, float, float]   # 上、下、左、右
    title: _FontSpec
    meta: _FontSpec                   # 标题下方的单位/日期行、落款
    headings: tuple[_FontSpec, _FontSpec, _FontSpec]
    body: _FontSpec
    table: _FontSpec
    line_spacing: float               # >= 6 视为固定磅值，否则视为倍数
    heading_indent: bool              # 一二级标题是否也首行缩进两字符
    numbering: str                    # chinese / decimal / none
    signature_at_end: bool            # 落款（单位、日期）放文末靠右
    meta_under_title: bool            # 单位/日期居中放在标题下方
    page_number_format: str           # "— {n} —" / "{n}"
    heading_space: tuple[tuple[float, float], ...] = field(
        default=((6, 0), (6, 0), (0, 0))
    )  # 各级标题段前/段后（磅）


TEMPLATES: dict[str, TemplateSpec] = {
    TEMPLATE_WORK_SUMMARY: TemplateSpec(
        key=TEMPLATE_WORK_SUMMARY,
        summary=(
            "公文式：A4，上 3.7cm/下 3.5cm/左 2.8cm/右 2.6cm；标题二号宋体加粗居中；"
            "一级标题黑体三号「一、」，二级楷体三号「（一）」，三级仿宋三号加粗「1.」；"
            "正文仿宋_GB2312 三号，首行缩进两字符，固定行距 28 磅；表格宋体五号带边框；"
            "页码「— 1 —」居中；文末单位、日期靠右落款。"
        ),
        margins_cm=(3.7, 3.5, 2.8, 2.6),
        title=_FontSpec("宋体", "Times New Roman", 22, bold=True),
        meta=_FontSpec("仿宋_GB2312", "Times New Roman", 16),
        headings=(
            _FontSpec("黑体", "Times New Roman", 16),
            _FontSpec("楷体_GB2312", "Times New Roman", 16),
            _FontSpec("仿宋_GB2312", "Times New Roman", 16, bold=True),
        ),
        body=_FontSpec("仿宋_GB2312", "Times New Roman", 16),
        table=_FontSpec("宋体", "Times New Roman", 10.5),
        line_spacing=28,
        heading_indent=True,
        numbering="chinese",
        signature_at_end=True,
        meta_under_title=False,
        page_number_format="— {n} —",
        heading_space=((0, 0), (0, 0), (0, 0)),
    ),
    TEMPLATE_FORMAL_REPORT: TemplateSpec(
        key=TEMPLATE_FORMAL_REPORT,
        summary=(
            "报告式：A4，上下 2.54cm/左右 3.17cm；标题二号黑体加粗居中，单位与日期居中列于标题下；"
            "一级标题黑体小三「1」，二级黑体四号「1.1」，三级黑体小四「1.1.1」；"
            "正文宋体小四，首行缩进两字符，1.5 倍行距；表格宋体五号带边框、表头加粗底纹；页码居中。"
        ),
        margins_cm=(2.54, 2.54, 3.17, 3.17),
        title=_FontSpec("黑体", "Times New Roman", 22, bold=True),
        meta=_FontSpec("宋体", "Times New Roman", 12),
        headings=(
            _FontSpec("黑体", "Times New Roman", 15, bold=True),
            _FontSpec("黑体", "Times New Roman", 14, bold=True),
            _FontSpec("黑体", "Times New Roman", 12, bold=True),
        ),
        body=_FontSpec("宋体", "Times New Roman", 12),
        table=_FontSpec("宋体", "Times New Roman", 10.5),
        line_spacing=1.5,
        heading_indent=False,
        numbering="decimal",
        signature_at_end=False,
        meta_under_title=True,
        page_number_format="{n}",
        heading_space=((12, 6), (6, 3), (3, 0)),
    ),
}

_TEMPLATE_ALIASES = {
    "工作总结": TEMPLATE_WORK_SUMMARY,
    "总结": TEMPLATE_WORK_SUMMARY,
    "工作总结模板": TEMPLATE_WORK_SUMMARY,
    "work-summary": TEMPLATE_WORK_SUMMARY,
    "work_summary": TEMPLATE_WORK_SUMMARY,
    "summary": TEMPLATE_WORK_SUMMARY,
    "公文": TEMPLATE_WORK_SUMMARY,
    "正式报告": TEMPLATE_FORMAL_REPORT,
    "报告": TEMPLATE_FORMAL_REPORT,
    "正式报告模板": TEMPLATE_FORMAL_REPORT,
    "formal-report": TEMPLATE_FORMAL_REPORT,
    "formal_report": TEMPLATE_FORMAL_REPORT,
    "report": TEMPLATE_FORMAL_REPORT,
    "分析报告": TEMPLATE_FORMAL_REPORT,
    "科研报告": TEMPLATE_FORMAL_REPORT,
}


def resolve_template(raw: Any) -> Optional[TemplateSpec]:
    """把模型传来的模板名（含常见别名、大小写、书名号）归一到内置模板；不认识返回 None。"""
    name = str(raw or "").strip().strip("「」《》“”\"' ").lower()
    if not name:
        return None
    for alias, key in _TEMPLATE_ALIASES.items():
        if name == alias.lower():
            return TEMPLATES[key]
    return None


def template_catalog() -> list[dict]:
    return [{"template": spec.key, "spec": spec.summary} for spec in TEMPLATES.values()]


# ---------------------------------------------------------------------------
# Markdown / 纯文本 → 块结构
# ---------------------------------------------------------------------------


@dataclass
class Block:
    kind: str                     # heading / paragraph / list / table / quote
    text: str = ""
    level: int = 0                # heading 1-3
    ordered: bool = False
    index: int = 0
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


_CN_NUM = "一二三四五六七八九十百"
_RE_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_RE_CN_H1 = re.compile(rf"^[{_CN_NUM}]+[、．.]\s*\S")
_RE_CN_H2 = re.compile(rf"^[（(][{_CN_NUM}]+[)）]\s*\S")
_RE_CN_H3 = re.compile(r"^\d{1,2}[．.、]\s*\S")
_RE_ORDERED = re.compile(r"^\s*(\d{1,3})[.)、]\s+(.+)$")
_RE_UNORDERED = re.compile(r"^\s*[-*+•]\s+(.+)$")
_RE_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_RE_NUMBER_PREFIX = re.compile(
    rf"^([{_CN_NUM}]+[、．.]|[（(][{_CN_NUM}]+[)）]|\d+(\.\d+)*[．.、]?|第[{_CN_NUM}\d]+[章节条部分])\s*"
)
_RE_END_PUNCT = re.compile(r"[。；;，,：:！!？?]$")


def _split_table_row(line: str) -> list[str]:
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", body)]


def _looks_like_cn_heading(line: str, pattern: re.Pattern) -> bool:
    return bool(pattern.match(line)) and len(line) <= 40 and not _RE_END_PUNCT.search(line)


def parse_blocks(text: str) -> list[Block]:
    """Markdown 标题为准；整篇没有 Markdown 标题时，按「一、/（一）/1.」公文惯例识别标题层级。

    段落内的换行按 Markdown 语义合并；空行分段。表格只认管道表（第二行为 --- 分隔）。
    """
    lines = str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    has_md_heading = any(_RE_MD_HEADING.match(line.strip()) for line in lines)
    blocks: list[Block] = []
    paragraph: list[str] = []
    in_code = False
    ordered_index = 0

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(Block("paragraph", " ".join(part.strip() for part in paragraph)))
            paragraph = []

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            paragraph.append(raw.rstrip())
            i += 1
            continue
        if not line:
            flush()
            ordered_index = 0
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", line):
            flush()
            i += 1
            continue
        # 管道表：当前行含 |，下一行是分隔行
        if "|" in line and i + 1 < len(lines) and _RE_TABLE_SEP.match(lines[i + 1]):
            flush()
            headers = _split_table_row(line)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                cells = _split_table_row(lines[i])
                cells = (cells + [""] * len(headers))[: len(headers)]
                rows.append(cells)
                i += 1
            blocks.append(Block("table", headers=headers, rows=rows))
            continue
        md = _RE_MD_HEADING.match(line)
        if md:
            flush()
            level = min(len(md.group(1)), 3)
            blocks.append(Block("heading", md.group(2).strip(), level=level))
            i += 1
            continue
        if not has_md_heading:
            if _looks_like_cn_heading(line, _RE_CN_H1):
                flush()
                blocks.append(Block("heading", line, level=1))
                i += 1
                continue
            if _looks_like_cn_heading(line, _RE_CN_H2):
                flush()
                blocks.append(Block("heading", line, level=2))
                i += 1
                continue
            if _looks_like_cn_heading(line, _RE_CN_H3) and len(line) <= 30:
                flush()
                blocks.append(Block("heading", line, level=3))
                i += 1
                continue
        if line.startswith(">"):
            flush()
            blocks.append(Block("quote", line.lstrip("> ").strip()))
            i += 1
            continue
        ordered = _RE_ORDERED.match(line)
        if ordered:
            flush()
            ordered_index += 1
            blocks.append(Block("list", ordered.group(2).strip(), ordered=True, index=ordered_index))
            i += 1
            continue
        unordered = _RE_UNORDERED.match(line)
        if unordered:
            flush()
            blocks.append(Block("list", unordered.group(1).strip(), ordered=False))
            i += 1
            continue
        paragraph.append(line)
        i += 1
    flush()
    return blocks


def pop_title(blocks: list[Block]) -> Optional[str]:
    """正文开头唯一的一级标题视为文档标题（避免标题在正文里重复出现一次）。"""
    if blocks and blocks[0].kind == "heading" and blocks[0].level == 1:
        others = [b for b in blocks[1:] if b.kind == "heading" and b.level == 1]
        if not others:
            return blocks.pop(0).text
    return None


def normalize_heading_levels(blocks: list[Block]) -> None:
    """标题用 # 而章节用 ## 时（常见 Markdown 习惯），把最高一级提升为一级，层级相对关系不变。"""
    levels = [b.level for b in blocks if b.kind == "heading"]
    if not levels:
        return
    shift = min(levels) - 1
    if shift <= 0:
        return
    for block in blocks:
        if block.kind == "heading":
            block.level = max(1, min(3, block.level - shift))


def outline_of(blocks: list[Block]) -> list[dict]:
    return [{"level": b.level, "text": b.text} for b in blocks if b.kind == "heading"]


# ---------------------------------------------------------------------------
# DOCX 渲染
# ---------------------------------------------------------------------------


def _qn(tag: str):
    from docx.oxml.ns import qn

    return qn(tag)


def _apply_rfonts(rpr, spec: _FontSpec) -> None:
    """写全 rFonts 的四个字体属性并清掉主题字体引用（否则 Word 仍按主题字体渲染）。"""
    from docx.oxml import OxmlElement

    rfonts = rpr.find(_qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(_qn(f"w:{attr}"), None)
    rfonts.set(_qn("w:ascii"), spec.ascii)
    rfonts.set(_qn("w:hAnsi"), spec.ascii)
    rfonts.set(_qn("w:eastAsia"), spec.east)
    rfonts.set(_qn("w:cs"), spec.ascii)


def _style_font(style, spec: _FontSpec, color_black: bool = True) -> None:
    from docx.shared import Pt, RGBColor

    style.font.name = spec.ascii
    style.font.size = Pt(spec.size)
    style.font.bold = spec.bold
    style.font.italic = False
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)
    _apply_rfonts(style.element.get_or_add_rPr(), spec)


def _run_font(run, spec: _FontSpec, bold: Optional[bool] = None, italic: bool = False) -> None:
    from docx.shared import Pt, RGBColor

    run.font.name = spec.ascii
    run.font.size = Pt(spec.size)
    run.font.bold = spec.bold if bold is None else bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    _apply_rfonts(run._element.get_or_add_rPr(), spec)


def _set_line_spacing(paragraph_format, spacing: float) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    if spacing >= 6:
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph_format.line_spacing = Pt(spacing)
    else:
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph_format.line_spacing = spacing


def _first_line_indent_chars(paragraph, chars: int, font_size: float) -> None:
    """首行缩进 N 字符：写 w:firstLineChars（Word 精确按字符）并给磅值兜底。"""
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    paragraph.paragraph_format.first_line_indent = Pt(font_size * chars)
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(_qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(_qn("w:firstLineChars"), str(chars * 100))


def _add_inline_runs(paragraph, text: str, spec: _FontSpec, base_bold: Optional[bool] = None) -> None:
    """支持 **加粗**、*斜体*、`代码`、[文字](链接) 的最小内联渲染。"""
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _run_font(paragraph.add_run(part[2:-2]), spec, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            _run_font(paragraph.add_run(part[1:-1]), spec, bold=base_bold)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _run_font(paragraph.add_run(part[1:-1]), spec, bold=base_bold, italic=True)
        else:
            _run_font(paragraph.add_run(part), spec, bold=base_bold)


def _add_page_number(paragraph, fmt: str, spec: _FontSpec) -> None:
    from docx.oxml import OxmlElement

    before, _, after = fmt.partition("{n}")
    if before:
        _run_font(paragraph.add_run(before), spec)
    run = paragraph.add_run()
    _run_font(run, spec)
    begin = OxmlElement("w:fldChar")
    begin.set(_qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(_qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(_qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(_qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)
    if after:
        _run_font(paragraph.add_run(after), spec)


def _shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement

    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(_qn("w:val"), "clear")
    shd.set(_qn("w:color"), "auto")
    shd.set(_qn("w:fill"), fill)
    tcpr.append(shd)


class _Numbering:
    """标题自动编号：标题文本自带「一、」「1.1」等前缀时不再重复加。"""

    def __init__(self, mode: str):
        self.mode = mode
        self.counters = [0, 0, 0]

    def prefix(self, level: int, text: str) -> str:
        if self.mode == "none" or _RE_NUMBER_PREFIX.match(text):
            return ""
        index = level - 1
        self.counters[index] += 1
        for deeper in range(index + 1, 3):
            self.counters[deeper] = 0
        if self.mode == "chinese":
            n = self.counters[index]
            if level == 1:
                return f"{_cn_number(n)}、"
            if level == 2:
                return f"（{_cn_number(n)}）"
            return f"{n}."
        parts = [str(c) for c in self.counters[: level] if c]
        return ".".join(parts) + "  "


def _cn_number(n: int) -> str:
    digits = "零一二三四五六七八九"
    if n < 10:
        return digits[n]
    if n < 20:
        return "十" + (digits[n - 10] if n > 10 else "")
    if n < 100:
        return digits[n // 10] + "十" + (digits[n % 10] if n % 10 else "")
    return str(n)


def render_docx(
    blocks: list[Block],
    template: TemplateSpec,
    *,
    title: str = "",
    organization: str = "",
    date: str = "",
) -> bytes:
    """把块结构按模板规格写成 DOCX 字节。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    document = Document()

    # 页面
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    top, bottom, left, right = template.margins_cm
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.left_margin, section.right_margin = Cm(left), Cm(right)

    # 样式：Normal 与 Heading 1-3 全部覆写为模板字体，保留 Heading 样式便于 Word 导航窗格 / PDF 书签
    normal = document.styles["Normal"]
    _style_font(normal, template.body)
    _set_line_spacing(normal.paragraph_format, template.line_spacing)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    for level in range(1, 4):
        style = document.styles[f"Heading {level}"]
        _style_font(style, template.headings[level - 1])
        _set_line_spacing(style.paragraph_format, template.line_spacing)
        before, after = template.heading_space[level - 1]
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # 标题块
    if title.strip():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(12 if template.meta_under_title else 18)
        _run_font(paragraph.add_run(title.strip()), template.title)
    if template.meta_under_title:
        for value in (organization, date):
            if value.strip():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run_font(paragraph.add_run(value.strip()), template.meta)
        if (organization.strip() or date.strip()):
            document.paragraphs[-1].paragraph_format.space_after = Pt(18)

    numbering = _Numbering(template.numbering)
    body_size = template.body.size
    for block in blocks:
        if block.kind == "heading":
            spec = template.headings[block.level - 1]
            paragraph = document.add_paragraph(style=f"Heading {block.level}")
            prefix = numbering.prefix(block.level, block.text)
            _add_inline_runs(paragraph, prefix + block.text, spec)
            if template.heading_indent:
                _first_line_indent_chars(paragraph, 2, spec.size)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            _first_line_indent_chars(paragraph, 2, body_size)
            _add_inline_runs(paragraph, block.text, template.body)
        elif block.kind == "list":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(body_size * 2)
            marker = f"{block.index}." if block.ordered else "•"
            _add_inline_runs(paragraph, f"{marker} {block.text}", template.body)
        elif block.kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(body_size * 2)
            _add_inline_runs(paragraph, block.text, template.body)
            for run in paragraph.runs:
                run.font.italic = True
        elif block.kind == "table":
            columns = max(len(block.headers), 1)
            table = document.add_table(rows=1 + len(block.rows), cols=columns)
            table.style = document.styles["Table Grid"]
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for column, header in enumerate(block.headers):
                cell = table.cell(0, column)
                paragraph = cell.paragraphs[0]  # 新建单元格自带一个空段落、无 run
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_line_spacing(paragraph.paragraph_format, 1.0)
                _add_inline_runs(paragraph, header, template.table, base_bold=True)
                _shade_cell(cell, "F2F2F2")
            for row_index, row in enumerate(block.rows, start=1):
                for column, value in enumerate(row[:columns]):
                    cell = table.cell(row_index, column)
                    paragraph = cell.paragraphs[0]
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    _set_line_spacing(paragraph.paragraph_format, 1.0)
                    _add_inline_runs(paragraph, value, template.table)
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)

    # 落款
    if template.signature_at_end and (organization.strip() or date.strip()):
        document.add_paragraph()
        for value in (organization, date):
            if value.strip():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                paragraph.paragraph_format.right_indent = Pt(body_size * 2)
                _run_font(paragraph.add_run(value.strip()), template.meta)

    # 页码
    footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer, template.page_number_format, _FontSpec("宋体", "Times New Roman", 14))

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


# ---------------------------------------------------------------------------
# DOCX → Markdown（导入 / 转换）
# ---------------------------------------------------------------------------


def _heading_level_of(style_name: str) -> int:
    name = str(style_name or "").strip().lower()
    match = re.match(r"^(heading|标题)\s*(\d)$", name)
    if match:
        return min(int(match.group(2)), 3)
    if name == "title":
        return 1
    return 0


def docx_to_markdown(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(data))
    lines: list[str] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level_of(paragraph.style.name if paragraph.style is not None else "")
            style_name = (paragraph.style.name if paragraph.style is not None else "").lower()
            if level:
                lines.append(f"{'#' * level} {text}")
            elif "list" in style_name:
                lines.append(f"- {text}")
            else:
                lines.append(text)
            lines.append("")
        elif tag == "tbl":
            table = Table(element, document)
            rows = [[cell.text.strip().replace("|", "\\|") for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [(row + [""] * width)[:width] for row in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + "|".join([" --- "] * width) + "|")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")
    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------
# 沙箱 LibreOffice 转换
# ---------------------------------------------------------------------------


async def soffice_convert(filename: str, data: bytes, target_ext: str) -> bytes:
    """在沙箱里用 LibreOffice 把版式文档转成 target_ext（pdf / docx），与文件预览同一条链路。"""
    from app.services.sandbox.office_convert import OfficeConvertError, convert_with_soffice

    target_ext = target_ext.lower().lstrip(".")
    try:
        return await convert_with_soffice(filename, data, target_ext, timeout_ms=_SOFFICE_TIMEOUT_MS)
    except OfficeConvertError as exc:
        raise ValueError(f"文档转换为 {target_ext.upper()} 失败（沙箱 LibreOffice）：{exc}") from exc


async def docx_to_pdf(filename: str, data: bytes) -> bytes:
    return await soffice_convert(filename, data, "pdf")


# ---------------------------------------------------------------------------
# 工具执行
# ---------------------------------------------------------------------------


def _safe_stem(raw: Any, fallback: str) -> str:
    name = os.path.basename(str(raw or "").replace("\\", "/").strip())
    name = re.sub(r'[\x00-\x1f<>:"|?*]', "_", name).strip(". ")
    stem = os.path.splitext(name)[0] if name else ""
    return (stem or fallback or "文档")[:150]


async def _save_deliverable(runtime: dict, filename: str, data: bytes, ext: str, origin: dict) -> dict:
    from app.services.files import user_file_service

    user_id = str(runtime.get("user_id") or "").strip()
    preview_only = bool(runtime.get("preview_only"))
    saved = await (
        user_file_service.save_preview_bytes(
            user_id, filename, data, run_id=str(runtime.get("run_id") or "") or None,
        )
        if preview_only
        else user_file_service.save_generated_bytes(
            user_id,
            filename,
            data,
            thread_id=str(runtime.get("thread_id") or "") or None,
            run_id=str(runtime.get("run_id") or "") or None,
        )
    )
    receipt = {
        "id": str(saved.get("id") or ""),
        "filename": str(saved.get("filename") or filename),
        "mime": str(saved.get("mime") or _MIMES.get(ext, "application/octet-stream")),
        "size": int(saved.get("size") or len(data)),
        "source": "generated",
        "versionNo": int(saved.get("versionNo") or 1),
        "deliverable": True,
        "previewOnly": preview_only,
        "origin": {"runId": str(runtime.get("run_id") or ""), **origin},
    }
    if not receipt["id"]:
        raise ValueError("文件保存未返回 file_id")
    return receipt


async def _load_source_text(runtime: dict, source_file_id: str) -> tuple[str, str]:
    """读「我的文件」里的 DOCX / MD / TXT，返回 (Markdown 正文, 原文件名)。"""
    from app.services.files import user_file_service

    user_id = str(runtime.get("user_id") or "").strip()
    row, data = await user_file_service.read_bytes(user_id, source_file_id)
    name = str(row.filename or "")
    ext = os.path.splitext(name)[1].lower()
    if ext == ".docx":
        return docx_to_markdown(data), name
    if ext == ".doc":
        converted = await soffice_convert(name, data, "docx")
        return docx_to_markdown(converted), name
    if ext in {".md", ".markdown", ".txt"}:
        return data.decode("utf-8-sig", errors="replace"), name
    raise ValueError("source_file_id 仅支持 DOCX、DOC、MD 或 TXT 正文文件；PDF/图片请先转成文字")


def _normalize_output_formats(raw: Any) -> list[str]:
    value = str(raw or "docx").strip().lower().replace("＋", "+").replace(" ", "")
    if value in {"docx+pdf", "pdf+docx", "both", "all"}:
        return ["docx", "pdf"]
    if value in {"docx", "word"}:
        return ["docx"]
    if value == "pdf":
        return ["pdf"]
    raise ValueError("output_format 仅支持 docx、pdf 或 docx+pdf")


async def typeset_document(args: dict, runtime: Optional[dict]) -> dict:
    """builtin.document_typeset：正文 + 内置模板 → 排版后的 DOCX / PDF。"""
    runtime = runtime or {}
    if not str(runtime.get("user_id") or "").strip():
        raise ValueError("缺少当前用户，无法保存排版文件")
    template = resolve_template(args.get("template"))
    if template is None:
        names = "、".join(f"「{item['template']}」" for item in template_catalog())
        raise ValueError(f"template 必须是内置模板之一：{names}")
    formats = _normalize_output_formats(args.get("output_format"))

    source_file_id = str(args.get("source_file_id") or "").strip()
    source_name = ""
    text = str(args.get("content") or "")
    if source_file_id:
        text, source_name = await _load_source_text(runtime, source_file_id)
    if not text.strip():
        raise ValueError("正文为空：请提供 content，或指定「我的文件」中的 DOCX/MD/TXT 作为 source_file_id")

    blocks = parse_blocks(text)
    title = str(args.get("title") or "").strip()
    embedded_title = pop_title(blocks)
    normalize_heading_levels(blocks)
    if not title:
        title = embedded_title or _safe_stem(source_name, "") or ""
    if not any(block.kind != "heading" for block in blocks):
        raise ValueError("正文只有标题没有内容，无法排版；请提供完整正文")
    organization = str(args.get("organization") or "").strip()
    date = str(args.get("date") or "").strip()
    docx_bytes = render_docx(blocks, template, title=title, organization=organization, date=date)

    stem = _safe_stem(args.get("filename"), title or _safe_stem(source_name, "排版文档") or "排版文档")
    origin = {"tool": "builtin.document_typeset", "template": template.key}
    if source_file_id:
        origin["sourceFileId"] = source_file_id
    receipts: list[dict] = []
    if "docx" in formats:
        receipts.append(await _save_deliverable(runtime, f"{stem}.docx", docx_bytes, "docx", origin))
    if "pdf" in formats:
        pdf_bytes = await docx_to_pdf(f"{stem}.docx", docx_bytes)
        receipts.append(await _save_deliverable(runtime, f"{stem}.pdf", pdf_bytes, "pdf", origin))

    outline = outline_of(blocks)
    tables = sum(1 for block in blocks if block.kind == "table")
    names = "、".join(f"《{item['filename']}》" for item in receipts)
    return {
        "message": f"已按「{template.key}」模板排版并保存 {names} 到我的文件，可在文件卡片直接预览。",
        "template": template.key,
        "template_spec": template.summary,
        "title": title,
        "outline": outline,
        "heading_counts": {
            "level1": sum(1 for item in outline if item["level"] == 1),
            "level2": sum(1 for item in outline if item["level"] == 2),
            "level3": sum(1 for item in outline if item["level"] == 3),
        },
        "table_count": tables,
        "files": receipts,
        "file": receipts[0],
        "_generated_file_receipt": receipts[0],
        "_generated_file_receipts": receipts,
    }


async def convert_document(args: dict, runtime: Optional[dict]) -> dict:
    """builtin.document_convert：我的文件里的源文件 + 目标格式 → 转换后的文件（不改排版内容）。"""
    from app.services.files import user_file_service

    runtime = runtime or {}
    user_id = str(runtime.get("user_id") or "").strip()
    source_file_id = str(args.get("source_file_id") or "").strip()
    if not user_id or not source_file_id:
        raise ValueError("缺少当前用户或 source_file_id，无法转换文件")
    target = str(args.get("target_format") or "").strip().lower().lstrip(".")
    if target == "word":
        target = "docx"
    if target == "markdown":
        target = "md"
    if target not in {"pdf", "docx", "md"}:
        raise ValueError("target_format 仅支持 pdf、docx 或 md")

    row, data = await user_file_service.read_bytes(user_id, source_file_id)
    source_name = str(row.filename or "源文件")
    ext = os.path.splitext(source_name)[1].lower().lstrip(".")
    if ext == target:
        raise ValueError(f"源文件已经是 {target.upper()} 格式，无需转换")
    stem = _safe_stem(args.get("filename"), _safe_stem(source_name, "转换结果"))

    if ext in {"docx", "doc"}:
        if target in {"pdf", "docx"}:
            output = await soffice_convert(source_name, data, target)
        else:
            source_docx = data if ext == "docx" else await soffice_convert(source_name, data, "docx")
            output = docx_to_markdown(source_docx).encode("utf-8")
    elif ext in {"md", "markdown", "txt"}:
        from app.services.skills.builtin_tools import _build_docx_bytes

        text = data.decode("utf-8-sig", errors="replace")
        blocks = parse_blocks(text)
        title = pop_title(blocks) or ""
        if target == "md":
            output = text.encode("utf-8")
        else:
            docx_bytes = _build_docx_bytes(title or stem, text if not title else _strip_first_heading(text))
            output = docx_bytes if target == "docx" else await docx_to_pdf(f"{stem}.docx", docx_bytes)
    else:
        raise ValueError("目前支持 DOCX/DOC → PDF/DOCX/MD，MD/TXT → DOCX/PDF；PPT、图片、扫描 PDF 暂不支持")

    receipt = await _save_deliverable(
        runtime,
        f"{stem}.{target}",
        output,
        target,
        {"tool": "builtin.document_convert", "sourceFileId": source_file_id, "sourceFormat": ext},
    )
    return {
        "message": f"已把《{source_name}》转换为 {target.upper()}，保存为《{receipt['filename']}》。",
        "source_format": ext,
        "target_format": target,
        "file": receipt,
        "_generated_file_receipt": receipt,
    }


def _strip_first_heading(text: str) -> str:
    lines = str(text or "").replace("\r\n", "\n").split("\n")
    for index, line in enumerate(lines):
        if not line.strip():
            continue
        if _RE_MD_HEADING.match(line.strip()):
            return "\n".join(lines[index + 1:])
        break
    return text


__all__ = [
    "TEMPLATES",
    "TEMPLATE_WORK_SUMMARY",
    "TEMPLATE_FORMAL_REPORT",
    "TEMPLATE_DEFAULT",
    "TemplateSpec",
    "Block",
    "resolve_template",
    "template_catalog",
    "parse_blocks",
    "pop_title",
    "normalize_heading_levels",
    "outline_of",
    "render_docx",
    "docx_to_markdown",
    "soffice_convert",
    "docx_to_pdf",
    "typeset_document",
    "convert_document",
]
