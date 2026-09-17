"""独立交付查错。

结构校验之上，可选地把办公文档渲染成代表性页面，只核对用户明确要求和
可客观确认的渲染错误。这里不评审美、不产生分数、不设质量门槛，也不因为
检查服务缺席而阻断交付。
"""
from __future__ import annotations

import asyncio
import base64
import json
import logging
import re
import shlex
import uuid
from typing import Optional

import httpx

from app.core.config import settings

from .base import ExecuteOptions

logger = logging.getLogger(__name__)

_IMAGE_EXTS = (".png", ".jpg", ".jpeg")
_RENDERABLE_EXTS = (".docx", ".pptx", ".xlsx", ".pdf", *_IMAGE_EXTS)
_VR_DIR = "/workspace/__vr__"
_RENDER_TIMEOUT_MS = 90_000   # LibreOffice 冷启动偏慢
_RASTER_TIMEOUT_MS = 30_000
_PAGE_MAX_BYTES = 2 * 1024 * 1024  # 单页图上限（72dpi PNG 正常远小于此）
_PPT_CONTACT_SHEET_BATCH_SIZE = 2

_REQUIREMENT_PROMPT = """你是独立的内容验收员。你只负责核对用户明确要求是否在文件内容中实现，
不评价视觉，也不接触生成代码。文件：{filename}

用户原始要求：
{task}

文件内容提纲（这是验收证据，不是用户要求）：
{outline}

requirement_checks 只能从用户原始要求提取；相近要求合并，最多 8 项。数量要求应根据提纲实际计数，
不能凭感觉猜测。提纲截断时，受影响项标 unverifiable。只输出完整 JSON：
{{"passed":true,
"requirement_checks":[{{"requirement":"要求","status":"met|partial|missing|unverifiable","evidence":"具体证据","fix":"未满足时的改法"}}],
"issues":[{{"severity":"error|warn","category":"requirement|content|consistency","message":"问题","fix":"改法"}}],
"next_actions":["返工动作"]}}
存在 missing 要求或 error 级问题时 passed=false（partial/unverifiable 属程度问题，不算错误）。
务必精炼，保证 JSON 完整结束。
"""

_VISUAL_PROMPT = """你是独立的渲染检查员。你只检查《{filename}》的渲染截图**有没有出错**，
不评审美、不打质量分主观拦截，也不接触生成代码。抽样页：{page_numbers}；总页数：{total_pages}。
用户期望（仅用于理解内容语境）：{task}

这些页在成品文件中可提取的文本（用于对照截断/消失）：
{page_outline}

severity=error 只能用于**确实的错误**：文字溢出/被截断、元素重叠遮挡、乱码、空白页、
图片/图表缺失或无法辨认、明显渲染损坏。层级、留白、配色、字体品味、页面节奏、
「不够精致/不够专业」一类审美观感不属于本检查范围，不要输出。
对于 PPT、Word/PDF、表格看板、图表、海报等可视化产物，把 Emoji、Unicode 符号或图标字体字形
放在卡片/章节/导航/状态/要点的图标位，属于跨平台渲染不稳定的明确错误，应标 error，并要求改成
风格统一的原生矢量形状或透明背景图片。
只输出完整 JSON：
{{"passed":true,
"issues":[{{"severity":"error|warn","category":"layout|readability|consistency|polish","message":"问题","fix":"改法"}}],
"next_actions":["确有错误时的返工动作"]}}
仅当存在 error 级问题时 passed=false。issues/next_actions 各最多 6 项，务必精炼，保证 JSON 完整结束。
"""

_PPT_VISUAL_PROMPT = """你是 PPT 渲染查错员，只检查《{filename}》的成品截图有没有明确错误。
抽样页：{page_numbers}；总页数：{total_pages}。用户目标：{task}

这些页在 PPTX 中可提取的完整文本（必须与截图对照，用于发现被截断、隐藏或没渲染的字）：
{page_outline}

先以 100% 视图逐页排查渲染硬伤。以下情况一律标 severity=error、passed=false：
- 大数字被意外拆行（例如 1700 显示成 170 + 0）、负号/百分号/单位错位；
- 标题、数字、单位、正文彼此重叠，或时间轴标签压住说明文字；
- 意外换行、文字截断、文字越出容器/画布、乱码、缺图或空白页。
必须特别放大检查 KPI 卡、时间轴、图表数据标签和页脚，并在 message 中写明具体页码与对象。

图片中出现素材库水印、网站角标、新闻字幕/截图 UI、价格标签、明显像素化或压缩马赛克时，
必须写明具体页码；影响内容辨识或用户明确要求时标 severity=error。
逐页检查图上文字的真实可读性：黑字压在深色/繁杂照片上、灰色副标题融进照片、标题跨过主体
或背景纹理导致字形边缘无法快速辨认，均属于明确 readability 问题；标题或关键副标题明显读不清时
标 severity=error。不能仅因页面存在半透明遮罩就假定对比足够。
如果出现方框/缺字，或疑似 Emoji 图标未被当前字体渲染，必须标 error。
不评价构图品味、高级感、模板感、卡片占比或页面节奏；这些由创作模型和用户决定。
只输出完整 JSON：
{{"passed":true,
"issues":[{{"severity":"error|warn","category":"layout|readability|rendering|imagery","message":"具体到页的错误","fix":"可直接执行的修改"}}],
"next_actions":["确有错误时的修动作"]}}
issues/next_actions 各最多 6 项，保证 JSON 完整结束。
"""

_OUTLINE_SCRIPT = r"""
import json, os, sys
p = sys.argv[1]
ext = os.path.splitext(p)[1].lower()
chunks, stats = [], {}
try:
    if ext == '.docx':
        from docx import Document
        d = Document(p)
        stats = {'paragraphs': len(d.paragraphs), 'tables': len(d.tables)}
        chunks.extend(x.text.strip() for x in d.paragraphs if x.text.strip())
        for ti, table in enumerate(d.tables, 1):
            chunks.append(f'[表格 {ti}]')
            for row in table.rows:
                line = ' | '.join(c.text.strip() for c in row.cells)
                if line.strip(' |'): chunks.append(line)
    elif ext == '.pptx':
        from pptx import Presentation
        d = Presentation(p)
        stats = {'slides': len(d.slides)}
        for i, slide in enumerate(d.slides, 1):
            chunks.append(f'[第 {i} 页]')
            for shape in slide.shapes:
                text = getattr(shape, 'text', '').strip()
                if text: chunks.append(text)
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        d = load_workbook(p, read_only=True, data_only=True)
        stats = {'sheets': d.sheetnames}
        remaining = 1800
        for ws in d.worksheets:
            chunks.append(f'[工作表 {ws.title}：{ws.max_row} 行 × {ws.max_column} 列]')
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals: chunks.append(' | '.join(vals))
                remaining -= len(vals)
                if remaining <= 0: break
            if remaining <= 0: break
    elif ext == '.pdf':
        from pypdf import PdfReader
        d = PdfReader(p)
        stats = {'pages': len(d.pages)}
        for i, page in enumerate(d.pages, 1):
            chunks.append(f'[第 {i} 页]')
            chunks.append((page.extract_text() or '').strip())
except Exception as exc:
    print(json.dumps({'outline': '', 'stats': stats, 'error': str(exc)[:300]}, ensure_ascii=False))
    raise SystemExit
text = '\n'.join(x for x in chunks if x)
limit = 16000
print(json.dumps({'outline': text[:limit], 'truncated': len(text) > limit, 'stats': stats}, ensure_ascii=False))
"""

_CONTACT_SHEET_SCRIPT = r"""
import json, os, sys
from PIL import Image, ImageDraw, ImageOps

workdir, names_json, output = sys.argv[1:4]
names = json.loads(names_json)
images = []
for name in names:
    path = os.path.join(workdir, name)
    with Image.open(path) as image:
        images.append((name, image.convert('RGB').copy()))
if not images:
    raise SystemExit(2)

cols = 1 if len(images) == 1 else 2
rows = (len(images) + cols - 1) // cols
cell_w, cell_h, label_h, gap = 960, 540, 36, 18
sheet = Image.new('RGB', (cols * cell_w + (cols + 1) * gap,
                          rows * (cell_h + label_h) + (rows + 1) * gap), '#E5E7EB')
draw = ImageDraw.Draw(sheet)
resample = getattr(Image, 'Resampling', Image).LANCZOS
for index, (name, image) in enumerate(images):
    row, col = divmod(index, cols)
    x = gap + col * (cell_w + gap)
    y = gap + row * (cell_h + label_h)
    thumb = ImageOps.contain(image, (cell_w, cell_h), method=resample)
    px = x + (cell_w - thumb.width) // 2
    py = y + (cell_h - thumb.height) // 2
    sheet.paste(thumb, (px, py))
    page = name.split('-')[-1].split('.')[0]
    draw.rectangle((x, y + cell_h, x + cell_w, y + cell_h + label_h), fill='#111827')
    draw.text((x + 12, y + cell_h + 9), f'PAGE {page}', fill='#FFFFFF')
sheet.save(output, 'JPEG', quality=88, optimize=True)
"""


def _is_renderable(name: str) -> bool:
    return name.lower().endswith(_RENDERABLE_EXTS)


def _sample_pages(total: int, limit: int) -> list[int]:
    """短文档全看；长文档均匀抽样且一定包含首页、末页。"""
    if total <= limit:
        return list(range(1, total + 1))
    if limit <= 1:
        return [1]
    return sorted({1 + round(i * (total - 1) / (limit - 1)) for i in range(limit)})


def _review_page_limit(filename: str, configured_limit: int) -> int:
    """常见 7–16 页 PPT 全量检查；普通文档仍尊重配置以控制延迟。"""
    return max(configured_limit, 16) if str(filename).lower().endswith(".pptx") else configured_limit


async def _extract_outline(
    sandbox, name: str, *, directory: str = "/workspace/outputs"
) -> dict:  # noqa: ANN001
    if name.lower().endswith((".png", ".jpg", ".jpeg")):
        return {"outline": "（独立图片，无可提取文本；以截图为准）", "truncated": False, "stats": {}}
    src = f"{directory.rstrip('/')}/{name}"
    cmd = f"python -c {shlex.quote(_OUTLINE_SCRIPT)} {shlex.quote(src)}"
    res = await sandbox.execute(cmd, ExecuteOptions(timeout_ms=30_000))
    try:
        data = json.loads((res.stdout or "").strip())
    except (ValueError, TypeError):
        return {"outline": "（内容提纲提取失败）", "truncated": True, "stats": {}}
    return data if isinstance(data, dict) else {"outline": "（内容提纲提取失败）"}


async def _render_pages(
    sandbox, name: str, *, directory: str = "/workspace/outputs"
) -> dict:  # noqa: ANN001
    """同沙箱内把产物转成代表性页面，返回图像、抽样页码和总页数。"""
    src = f"{directory.rstrip('/')}/{name}"
    if name.lower().endswith((".png", ".jpg", ".jpeg")):
        results = await sandbox.read_files([src])
        if not results or not results[0].ok or not results[0].data:
            return {"pages": [], "page_numbers": [], "total_pages": 1}
        mime = "image/png" if name.lower().endswith(".png") else "image/jpeg"
        return {
            "pages": [{"data": results[0].data, "mime": mime}],
            "page_numbers": [1],
            "total_pages": 1,
        }
    token = uuid.uuid4().hex[:10]
    stem = name.rsplit(".", 1)[0]
    lower = name.lower()
    workdir = f"{_VR_DIR}/{token}"
    if lower.endswith(".pdf"):
        prep = f"mkdir -p {shlex.quote(workdir)} && cp {shlex.quote(src)} {shlex.quote(workdir + '/doc.pdf')}"
    else:
        # to-pdf 输出文件名 = 输入 stem + .pdf
        prep = (
            f"mkdir -p {shlex.quote(workdir)} && to-pdf {shlex.quote(src)} {shlex.quote(workdir)} "
            f"&& mv {shlex.quote(workdir + '/' + stem + '.pdf')} {shlex.quote(workdir + '/doc.pdf')}"
        )
    res = await sandbox.execute(prep, ExecuteOptions(timeout_ms=_RENDER_TIMEOUT_MS))
    if not res.ok:
        logger.warning("视觉审查渲染 PDF 失败 %s: %s", name, (res.stderr or res.stdout)[:200])
        return {"pages": [], "page_numbers": [], "total_pages": 0}
    page_count_cmd = (
        "from pypdf import PdfReader;import sys;print(len(PdfReader(sys.argv[1]).pages))"
    )
    count_res = await sandbox.execute(
        f"python -c {shlex.quote(page_count_cmd)} {shlex.quote(workdir + '/doc.pdf')}",
        ExecuteOptions(timeout_ms=10_000),
    )
    try:
        total_pages = max(1, int((count_res.stdout or "").strip()))
    except (TypeError, ValueError):
        total_pages = int(settings.SANDBOX_VISUAL_REVIEW_MAX_PAGES)
    configured_limit = int(settings.SANDBOX_VISUAL_REVIEW_MAX_PAGES)
    page_limit = _review_page_limit(name, configured_limit)
    page_numbers = _sample_pages(total_pages, page_limit)
    for page_no in page_numbers:
        dpi = 96 if lower.endswith(".pptx") else 72
        raster = (
            f"pdftoppm -png -r {dpi} -f {page_no} -l {page_no} -singlefile "
            f"{shlex.quote(workdir + '/doc.pdf')} {shlex.quote(workdir + f'/page-{page_no}') }"
        )
        res = await sandbox.execute(raster, ExecuteOptions(timeout_ms=_RASTER_TIMEOUT_MS))
        if not res.ok:
            logger.warning("视觉审查转页图失败 %s p%s: %s", name, page_no, (res.stderr or res.stdout)[:200])
    lister = (
        "import os,json;d=" + repr(workdir) + ";"
        "print(json.dumps(sorted((n for n in os.listdir(d) if n.startswith('page-') and n.endswith('.png')),"
        "key=lambda n:int(n.split('-')[1].split('.')[0]))))"
    )
    res = await sandbox.execute(f"python -c {shlex.quote(lister)}", ExecuteOptions(timeout_ms=10_000))
    try:
        names = json.loads((res.stdout or "").strip() or "[]")
    except (ValueError, TypeError):
        return {"pages": [], "page_numbers": [], "total_pages": total_pages}
    if not names:
        return {"pages": [], "page_numbers": [], "total_pages": total_pages}
    # 单次视觉请求只接收 1 张图，但不能把 10–16 页全塞进一张纵向联系表。
    # API 下采样后，这种大图里每页只剩约 400×225，文字重叠和大数字截断很容易漏检。
    # 最多 2 页一张联系表，由 _ask_vision 逐批调用；在常见端点
    # 的图像缩放上限下，仍能保住 KPI、时间轴标签等小字的可读性。
    batches = []
    for batch_start in range(0, len(names), _PPT_CONTACT_SHEET_BATCH_SIZE):
        batch_names = names[batch_start:batch_start + _PPT_CONTACT_SHEET_BATCH_SIZE]
        batch_no = batch_start // _PPT_CONTACT_SHEET_BATCH_SIZE + 1
        sheet_path = f"{workdir}/contact-sheet-{batch_no}.jpg"
        sheet_cmd = (
            f"python -c {shlex.quote(_CONTACT_SHEET_SCRIPT)} {shlex.quote(workdir)} "
            f"{shlex.quote(json.dumps(batch_names))} {shlex.quote(sheet_path)}"
        )
        sheet_res = await sandbox.execute(sheet_cmd, ExecuteOptions(timeout_ms=20_000))
        if not sheet_res.ok:
            logger.warning(
                "视觉审查联系表生成失败 %s batch=%s: %s",
                name, batch_no, (sheet_res.stderr or sheet_res.stdout)[:200],
            )
            batches = []
            break
        results = await sandbox.read_files([sheet_path])
        if (
            not results or not results[0].ok or not results[0].data
            or len(results[0].data) > _PAGE_MAX_BYTES
        ):
            logger.warning("视觉审查联系表读回失败 %s batch=%s", name, batch_no)
            batches = []
            break
        batch_pages = [int(item.split("-")[1].split(".")[0]) for item in batch_names]
        batches.append({
            "data": results[0].data,
            "mime": "image/jpeg",
            "page_numbers": batch_pages,
        })
    if batches:
        return {"pages": batches, "page_numbers": page_numbers, "total_pages": total_pages}
    logger.warning("视觉审查联系表批处理失败 %s，降级为首页", name)
    # Pillow 或联系表生成异常时只发首张，至少保留硬伤检查且不再次触发图片数量限制。
    results = await sandbox.read_files([f"{workdir}/{names[0]}"])
    blobs = [
        {"data": r.data, "mime": "image/png"}
        for r in results if r.ok and r.data and len(r.data) <= _PAGE_MAX_BYTES
    ]
    return {
        "pages": [dict(item, page_numbers=page_numbers[:1]) for item in blobs],
        "page_numbers": page_numbers,
        "total_pages": total_pages,
    }


def _parse_verdict(text: str) -> Optional[dict]:
    """容错解析模型 JSON（可能包 ```json 围栏或前后缀文字）。"""
    m = re.search(r"\{.*\}", text or "", re.S)
    if not m:
        return None
    try:
        data = json.loads(m.group(0))
    except (ValueError, TypeError):
        return None
    if not isinstance(data, dict) or "passed" not in data:
        return None
    issues = []
    for item in data.get("issues") or []:
        if not isinstance(item, dict):
            continue
        issues.append({
            "severity": "error" if str(item.get("severity")) == "error" else "warn",
            "category": str(item.get("category") or "polish")[:40],
            "message": str(item.get("message") or "")[:300],
            "fix": str(item.get("fix") or "")[:300],
        })
    checks = []
    for item in data.get("requirement_checks") or []:
        if not isinstance(item, dict):
            continue
        status = str(item.get("status") or "unverifiable")
        if status not in ("met", "partial", "missing", "unverifiable"):
            status = "unverifiable"
        checks.append({
            "requirement": str(item.get("requirement") or "")[:240],
            "status": status,
            "evidence": str(item.get("evidence") or "")[:300],
            "fix": str(item.get("fix") or "")[:300],
        })
    verdict = {
        "passed": bool(data.get("passed")),
        "requirement_checks": checks,
        "issues": issues[:12],
        "next_actions": [str(x)[:300] for x in (data.get("next_actions") or []) if str(x).strip()][:8],
    }
    # 审查只查「错」，不评「质量」（2026-07-15 用户拍板）：质量不满意由用户在对话里
    # 继续迭代，不由门禁替用户做主。拦截条件只有两类确实的错误——
    # ① error 级问题（溢出/遮挡/乱码/空白/缺图等渲染硬伤）；② 用户明确要求 missing（没做）。
    # partial/unverifiable 属程度问题放行。本模块不产生或保存任何质量分。
    has_missing = any(x["status"] == "missing" for x in checks)
    has_error = any(x["severity"] == "error" for x in issues)
    verdict["passed"] = not has_error and not has_missing
    return verdict


def _parse_requirement_verdict(text: str) -> Optional[dict]:
    m = re.search(r"\{.*\}", text or "", re.S)
    if not m:
        return None
    try:
        data = json.loads(m.group(0))
    except (ValueError, TypeError):
        return None
    if not isinstance(data, dict) or not isinstance(data.get("requirement_checks"), list):
        return None
    checks = []
    for item in data["requirement_checks"][:8]:
        if not isinstance(item, dict):
            continue
        status = str(item.get("status") or "unverifiable")
        if status not in ("met", "partial", "missing", "unverifiable"):
            status = "unverifiable"
        checks.append({
            "requirement": str(item.get("requirement") or "")[:240],
            "status": status,
            "evidence": str(item.get("evidence") or "")[:300],
            "fix": str(item.get("fix") or "")[:300],
        })
    return {
        "passed": bool(data.get("passed")),
        "requirement_checks": checks,
        "issues": [x for x in (data.get("issues") or []) if isinstance(x, dict)][:8],
        "next_actions": [str(x)[:300] for x in (data.get("next_actions") or []) if str(x).strip()][:6],
    }


def _parse_visual_verdict(text: str) -> Optional[dict]:
    m = re.search(r"\{.*\}", text or "", re.S)
    if not m:
        return None
    try:
        data = json.loads(m.group(0))
    except (ValueError, TypeError):
        return None
    if not isinstance(data, dict) or "passed" not in data:
        return None
    issues = []
    for item in data.get("issues") or []:
        if not isinstance(item, dict):
            continue
        issues.append({
            "severity": "error" if str(item.get("severity")) == "error" else "warn",
            "category": str(item.get("category") or "rendering")[:40],
            "message": str(item.get("message") or "")[:300],
            "fix": str(item.get("fix") or "")[:300],
        })
    has_error = any(item["severity"] == "error" for item in issues)
    return {
        "passed": not has_error,
        "issues": issues[:8],
        "next_actions": [str(x)[:300] for x in (data.get("next_actions") or []) if str(x).strip()][:6],
    }


def _unique_values(values: list, limit: int) -> list:
    output = []
    seen = set()
    for value in values:
        marker = (
            json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, dict) else str(value)
        )
        if marker and marker not in seen:
            output.append(value)
            seen.add(marker)
        if len(output) >= limit:
            break
    return output


def _merge_visual_batch_verdicts(items: list[dict]) -> Optional[dict]:
    """Merge independent page-batch error reports."""
    if not items:
        return None
    merged = dict(items[0])
    merged["passed"] = all(bool(item.get("passed")) for item in items)
    merged["issues"] = _unique_values(
        [value for item in items for value in (item.get("issues") or [])], 12,
    )
    merged["next_actions"] = _unique_values(
        [value for item in items for value in (item.get("next_actions") or [])], 8,
    )
    return merged


def _outline_for_pages(outline_text: str, page_numbers: list[int]) -> str:
    """Select the PPT/PDF outline sections that correspond to one image batch."""
    wanted = {int(value) for value in page_numbers}
    sections = re.findall(
        r"(?ms)^\[第\s*(\d+)\s*页\]\s*\n?(.*?)(?=^\[第\s*\d+\s*页\]|\Z)",
        str(outline_text or ""),
    )
    selected = [f"[第 {page} 页]\n{body.strip()}" for page, body in sections if int(page) in wanted]
    if selected:
        return "\n".join(selected)[:5000]
    return (str(outline_text or "").strip() or "（无可提取文本）")[:2500]


async def _ask_vision(
    rendered: dict, *, filename: str, task_brief: str, outline: dict, newapi_key: str
) -> Optional[dict]:
    """调用独立视觉模型（复用 OCR multimodal_model 配置）。不可用/失败返回 None。"""
    from app.services.platform import platform_config_service as cfg

    conf = await cfg.get_ocr_config()
    model = str(conf.get("model") or "").strip()
    if not model:
        return None
    vision_base = str(conf.get("visionBaseUrl") or "").strip().rstrip("/")
    if vision_base:
        base_url, api_key = vision_base, str(conf.get("visionApiKey") or "").strip()
    else:
        base_url, api_key = settings.NEWAPI_BASE_URL.rstrip("/"), newapi_key
    if not api_key:
        return None

    task = (task_brief or "（未提供明确任务，只查渲染错误）")[:2400]
    outline_text = ((outline.get("outline") or "（无可用提纲）")[:16000]
                    + ("\n[提纲因长度已截断]" if outline.get("truncated") else ""))
    requirement_content: list = [{"type": "text", "text": _REQUIREMENT_PROMPT.format(
        filename=filename, task=task, outline=outline_text,
    )}]
    visual_prompt = _PPT_VISUAL_PROMPT if filename.lower().endswith(".pptx") else _VISUAL_PROMPT
    visual_batches: list[list] = []
    for page in rendered.get("pages") or []:
        batch_numbers = page.get("page_numbers") or rendered.get("page_numbers") or []
        visual_content: list = [{"type": "text", "text": visual_prompt.format(
            filename=filename,
            page_numbers="、".join(str(x) for x in batch_numbers) or "无",
            total_pages=rendered.get("total_pages") or "未知",
            task=task,
            page_outline=_outline_for_pages(outline_text, batch_numbers),
        )}]
        image_item = {
            "type": "image_url",
            "image_url": {"url": (
                f"data:{page.get('mime') or 'image/png'};base64,"
                f"{base64.b64encode(page['data']).decode()}"
            )},
        }
        visual_content.append(image_item)
        visual_batches.append(visual_content)
        # 独立图片没有可提取的正文提纲，需求验收员必须看到图片本身才能核对文字/对象要求。
        if filename.lower().endswith((".png", ".jpg", ".jpeg")):
            requirement_content.append(image_item)

    async def ask(content: list, parser, label: str) -> Optional[dict]:  # noqa: ANN001
        from app.services.agent_harness import model_usage_audit
        from app.services.chat.tools.base import current_tool_context

        tool_context = current_tool_context()
        run_id = str(tool_context.run_id or "") if tool_context is not None else ""
        thread_id = str(tool_context.thread_id or "") if tool_context is not None else ""
        parent_tool_call_id = str(tool_context.call_id or "") if tool_context is not None else ""
        logical = None
        if run_id:
            logical = await model_usage_audit.begin_logical_call(
                run_id=run_id,
                thread_id=thread_id,
                parent_tool_call_id=parent_tool_call_id,
                model=model,
                transport="chat_completions",
                purpose="tool_internal",
                purpose_detail="sandbox_visual_review",
                scope_key="tool_internal:sandbox_visual_review",
                provider_api_key=api_key,
            )
        else:
            logger.warning(
                "model_usage_orphan purpose=tool_internal purpose_detail=sandbox_visual_review "
                "reason=missing_run_id"
            )
        async with httpx.AsyncClient(timeout=90) as client:
            previous_attempt_id = ""
            for attempt_index in range(2):
                wire_payload = {
                    "model": model,
                    "messages": [
                        {"role": "system", "content": "你是独立、精炼的文件查错员。不评分，只输出完整 JSON。"},
                        {"role": "user", "content": content},
                    ],
                    "temperature": 0.1,
                    # 当前平台默认 glm-4v-flash 的上限为 1024；结构要求刻意保持精炼，
                    # 兼容其它 OpenAI-compatible 视觉端点。
                    "max_tokens": 1024,
                }
                audit_attempt = (
                    await model_usage_audit.begin_attempt(
                        logical,
                        wire_payload=wire_payload,
                        attempt_kind="initial" if attempt_index == 0 else "retry",
                        retry_of_attempt_id=previous_attempt_id,
                        legacy_compatible=False,
                    )
                    if logical is not None else None
                )
                if audit_attempt is not None:
                    previous_attempt_id = audit_attempt.attempt_id
                try:
                    resp = await client.post(
                        f"{base_url}/chat/completions",
                        headers={"Authorization": f"Bearer {api_key}"},
                        json=wire_payload,
                    )
                except asyncio.CancelledError:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="cancelled",
                        provider_event_seen=False,
                        unknown_provider_charge=True,
                        committed=False,
                    )
                    await model_usage_audit.finish_logical_call(
                        logical, terminal_status="cancelled", committed=False,
                    )
                    raise
                except httpx.TransportError as exc:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="failed",
                        provider_event_seen=False,
                        unknown_provider_charge=True,
                        error_code=type(exc).__name__,
                        committed=False,
                    )
                    logger.warning(
                        "独立%s审查连接失败 %s attempt=%s error=%s",
                        label, filename, attempt_index + 1, type(exc).__name__,
                    )
                    if attempt_index == 0:
                        continue
                    break
                except Exception as exc:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="failed",
                        provider_event_seen=False,
                        unknown_provider_charge=True,
                        error_code=type(exc).__name__,
                        committed=False,
                    )
                    await model_usage_audit.finish_logical_call(
                        logical, terminal_status="failed", committed=False,
                    )
                    raise
                try:
                    data = resp.json()
                    data_error: Optional[Exception] = None
                except Exception as exc:
                    data = {}
                    data_error = exc
                if resp.status_code >= 400:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="http_error",
                        usage=model_usage_audit.provider_usage_from_response(data),
                        response_id=model_usage_audit.provider_response_id(data),
                        provider_event_seen=True,
                        terminal_seen=True,
                        http_status=resp.status_code,
                        committed=False,
                    )
                    logger.warning(
                        "独立%s审查请求失败 %s attempt=%s status=%s body=%s",
                        label, filename, attempt_index + 1, resp.status_code, resp.text[:500],
                    )
                    if attempt_index == 0 and (
                        resp.status_code in {408, 425} or resp.status_code >= 500
                    ):
                        continue
                    await model_usage_audit.finish_logical_call(
                        logical, terminal_status="failed", committed=False,
                    )
                    return None
                if data_error is not None:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="invalid_response",
                        provider_event_seen=True,
                        terminal_seen=True,
                        http_status=resp.status_code,
                        error_code=type(data_error).__name__,
                        committed=False,
                    )
                    await model_usage_audit.finish_logical_call(
                        logical, terminal_status="failed", committed=False,
                    )
                    return None
                text = str(((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "")
                verdict = parser(text)
                if verdict is not None:
                    await model_usage_audit.finish_attempt(
                        audit_attempt,
                        terminal_status="completed",
                        usage=model_usage_audit.provider_usage_from_response(data),
                        response_id=model_usage_audit.provider_response_id(data),
                        provider_event_seen=True,
                        terminal_seen=True,
                        http_status=resp.status_code,
                        committed=True,
                    )
                    await model_usage_audit.finish_logical_call(
                        logical,
                        terminal_status="completed",
                        selected_attempt_id=(
                            audit_attempt.attempt_id if audit_attempt is not None else ""
                        ),
                        committed=True,
                    )
                    return verdict
                await model_usage_audit.finish_attempt(
                    audit_attempt,
                    terminal_status="invalid_response",
                    usage=model_usage_audit.provider_usage_from_response(data),
                    response_id=model_usage_audit.provider_response_id(data),
                    provider_event_seen=True,
                    terminal_seen=True,
                    http_status=resp.status_code,
                    committed=False,
                )
                logger.warning(
                    "独立%s审查返回不可解析 %s attempt=%s content=%s",
                    label, filename, attempt_index + 1, text[:500],
                )
                await model_usage_audit.finish_logical_call(
                    logical, terminal_status="failed", committed=False,
                )
                return None
        await model_usage_audit.finish_logical_call(
            logical, terminal_status="failed", committed=False,
        )
        return None

    try:
        requirement = await ask(requirement_content, _parse_requirement_verdict, "需求")
        visual_results = []
        for index, content in enumerate(visual_batches, 1):
            visual_results.append(await ask(
                content, _parse_visual_verdict,
                f"视觉第 {index}/{len(visual_batches)} 批",
            ))
        if requirement is None or not visual_results or any(item is None for item in visual_results):
            return None
        visual_items = [item for item in visual_results if item is not None]
        visual = _merge_visual_batch_verdicts(visual_items)
        if visual is None:
            return None
        raw = {
            "passed": bool(requirement["passed"] and visual["passed"]),
            "requirement_checks": requirement["requirement_checks"],
            "issues": [*(requirement["issues"] or []), *(visual["issues"] or [])],
            "next_actions": [*(requirement["next_actions"] or []), *(visual["next_actions"] or [])],
        }
        return _parse_verdict(json.dumps(raw, ensure_ascii=False))
    except Exception as e:  # noqa: BLE001
        logger.warning("独立文件查错调用失败 %s: %s", filename, e)
        return None


def _merge_into_review(entry: dict, verdict: Optional[dict], *, required: Optional[bool] = None) -> None:
    """把客观查错结果并进结构 review；检查缺席不影响交付。"""
    review = entry.setdefault("review", {"status": "passed", "summary": ""})
    prefix = (review.get("summary") or "").strip()
    if verdict is None:
        # 审查服务异常/不可用**不得触发返工或拦截**（2026-07-15 用户拍板）：结构硬校验
        # 已经保证文件能打开、格式正确；渲染检查缺席只如实记录，不惩罚产物。
        note = "已完成结构校验；独立渲染检查未完成（不影响交付）"
        review["summary"] = f"{prefix}；{note}" if prefix else note
        return
    review["quality"] = verdict
    if verdict["passed"]:
        note = "独立渲染检查未发现明确错误"
        if verdict["issues"]:
            note += "（提醒：" + "；".join(i["message"] for i in verdict["issues"][:2]) + "）"
        review["summary"] = f"{prefix}；{note}" if prefix else note
        return
    review["status"] = "failed"
    # 查错语义（2026-07-15 拍板）：只报确实的错误（error 级问题 + missing 要求）；
    # partial 等程度问题不进「错误」清单——那属于质量迭代，交给用户
    actions = verdict.get("next_actions") or [i.get("fix") for i in verdict["issues"] if i.get("fix")]
    problems = "；".join(
        i["message"] for i in verdict["issues"][:4] if i.get("severity") == "error"
    )
    unmet = [x for x in verdict.get("requirement_checks") or [] if x.get("status") == "missing"]
    if unmet:
        problems = (problems + "；" if problems else "") + "要求未实现：" + "；".join(
            x.get("requirement") or "未命名要求" for x in unmet[:3]
        )
    note = f"检查发现需修正的错误：{problems or '产物存在错误'}"
    if actions:
        note += "；返工：" + "；".join(str(x) for x in actions[:3])
    review["summary"] = f"{prefix}；{note}" if prefix else note


async def visual_review_outputs(
    sandbox, outputs: list, *, task_brief: str = "", newapi_key: str = "",
    directory: str = "/workspace/outputs",
) -> Optional[dict]:  # noqa: ANN001
    """对硬校验通过的文档/图片做可选的需求与渲染查错。

    不产生评分；检查缺席和超出 MAX_FILES 都不拦截交付。
    """
    candidates_all = [
        o for o in outputs
        if _is_renderable(o.get("name") or "")
        and (o.get("review") or {}).get("status") in ("passed", "warning")
    ]
    # 同轮既有文档又有图片时，图片通常是 matplotlib 等生成后嵌入文档的中间素材，
    # 与文件区“只保存文档、不刷屏保存中间图”保持同一口径；只有图片产物时才逐张验收。
    if any(not str(o.get("name") or "").lower().endswith(_IMAGE_EXTS) for o in candidates_all):
        candidates_all = [
            o for o in candidates_all
            if not str(o.get("name") or "").lower().endswith(_IMAGE_EXTS)
        ]
    limit = int(settings.SANDBOX_VISUAL_REVIEW_MAX_FILES)
    candidates = candidates_all[:limit]
    overflow = candidates_all[limit:]
    if not candidates_all:
        return None
    worst = "passed"
    for o in candidates:
        rendered = await _render_pages(sandbox, o["name"], directory=directory)
        outline = await _extract_outline(sandbox, o["name"], directory=directory)
        verdict = await _ask_vision(
            rendered, filename=o["name"], task_brief=task_brief,
            outline=outline, newapi_key=newapi_key,
        ) if rendered.get("pages") else None
        _merge_into_review(o, verdict)
        status = (o.get("review") or {}).get("status")
        if status == "failed":
            worst = "failed"
        elif status == "warning" and worst == "passed":
            # 分级交付档：可交付但有差距——总体状态如实降为 warning（照常放行）
            worst = "warning"
        # verdict is None：条目已按「不影响交付」保持 passed，总体不得因此变成 warning。
    for o in overflow:
        # 超出审查能力上限=审查缺席，不是产物的错：不拦截，如实记录即可
        _merge_into_review(o, None)
        o["review"]["summary"] += f"；单轮可视产物超过 {limit} 个，本文件未做渲染检查"
        if worst == "passed":
            worst = "warning"
    items = [
        {
            "name": str(o.get("name") or ""),
            "status": str((o.get("review") or {}).get("status") or "unknown"),
            "message": str((o.get("review") or {}).get("summary") or ""),
            "quality": (o.get("review") or {}).get("quality"),
        }
        for o in candidates_all
    ]
    summary = "；".join(
        str(item.get("message") or "").strip()
        for item in items
        if str(item.get("message") or "").strip()
    )[:800]
    return {
        "status": worst,
        "summary": summary,
        "items": items,
    }
