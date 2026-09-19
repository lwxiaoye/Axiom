"""统一文档解析（ADR-040）：会话上传文件抽取为文本，供主对话作为上下文注入。

支持：纯文本/代码/csv/json、pdf（pypdf 提文字层；无文字层的扫描件逐页渲染成 PNG
走自建 OCR 端点、没填端点则走视觉模型，见 _ocr_pdf_scanned）、docx（python-docx）、pptx（python-pptx）、
图片（走 OCR 策略）。Word/PDF/PPT 中夹带的照片、截图和图表会额外走视觉识别。
OCR 策略读 `agent_platform_config.ocr`：custom_endpoint（自建端点）或 multimodal_model
（默认；视觉模型自动取平台对话模型连接，见 _resolve_vision_runtime）。

主对话里的图片附件**不走这里**：对话模型是多模态的，像素以 image 内容块直接进模型
（2026-09-19 用户拍板「OCR 不需要，我们用的是多模态模型」）。这里只服务知识库/文档解析/
纯文本模型这些确实需要把图变成文字的路径。
"""
import asyncio
import base64
import hashlib
import io
import logging
import zipfile
from typing import Any, Callable, Optional

import httpx

from app.core.model_endpoint import get_model_base_url
from app.core.config import settings
from app.services.platform import platform_config_service as cfg
from app.services.platform import zip_guard

logger = logging.getLogger(__name__)


async def _audited_document_model_post(
    client: httpx.AsyncClient,
    url: str,
    *,
    model: str,
    purpose_detail: str,
    provider_api_key: str,
    headers: dict,
    wire_payload: dict,
    response_has_output: Callable[[dict], bool],
    json_payload: Optional[dict] = None,
    files: Optional[dict] = None,
    audit_context: Optional[dict] = None,
) -> tuple[httpx.Response, dict]:
    """Send one OCR/vision Provider attempt and attribute it to its real lifecycle Run."""
    from app.services.agent_harness import model_usage_audit
    from app.services.chat.tools.base import current_tool_context

    explicit = dict(audit_context or {})
    tool_context = current_tool_context()
    run_id = str(
        explicit.get("run_id")
        or (tool_context.run_id if tool_context is not None else "")
        or ""
    )
    thread_id = str(
        explicit.get("thread_id")
        or (tool_context.thread_id if tool_context is not None else "")
        or ""
    )
    root_run_id = str(explicit.get("root_run_id") or "")
    parent_tool_call_id = str(
        explicit.get("parent_tool_call_id")
        or (tool_context.call_id if tool_context is not None else "")
        or ""
    )
    parent_logical_call_id = str(
        explicit.get("parent_logical_call_id")
        or (tool_context.parent_logical_call_id if tool_context is not None else "")
        or ""
    )
    execution_segment = str(
        explicit.get("execution_segment")
        or (tool_context.execution_segment if tool_context is not None else "")
        or ""
    )
    logical = None
    if run_id:
        logical = await model_usage_audit.begin_logical_call(
            run_id=run_id,
            root_run_id=root_run_id,
            thread_id=thread_id,
            parent_logical_call_id=parent_logical_call_id,
            parent_tool_call_id=parent_tool_call_id,
            model=model,
            transport=("chat_completions" if json_payload is not None else "custom_ocr"),
            purpose="tool_internal",
            purpose_detail=purpose_detail,
            scope_key=f"tool_internal:{purpose_detail}",
            provider_api_key=provider_api_key,
        )
    else:
        logger.warning(
            "model_usage_orphan purpose=tool_internal purpose_detail=%s reason=missing_run_id",
            purpose_detail,
        )
    attempt = (
        await model_usage_audit.begin_attempt(
            logical,
            wire_payload=wire_payload,
            attempt_kind="initial",
            execution_segment=execution_segment,
            legacy_compatible=False,
        )
        if logical is not None else None
    )
    try:
        request_kwargs: dict[str, Any] = {"headers": headers}
        if json_payload is not None:
            request_kwargs["json"] = json_payload
        if files is not None:
            request_kwargs["files"] = files
        response = await client.post(url, **request_kwargs)
    except asyncio.CancelledError:
        await model_usage_audit.finish_attempt(
            attempt,
            terminal_status="cancelled",
            provider_event_seen=False,
            unknown_provider_charge=True,
            committed=False,
        )
        await model_usage_audit.finish_logical_call(
            logical, terminal_status="cancelled", committed=False,
        )
        raise
    except Exception as exc:
        await model_usage_audit.finish_attempt(
            attempt,
            terminal_status="failed",
            provider_event_seen=False,
            error_code=type(exc).__name__,
            committed=False,
        )
        await model_usage_audit.finish_logical_call(
            logical, terminal_status="failed", committed=False,
        )
        raise

    try:
        data = response.json()
    except Exception as exc:
        await model_usage_audit.finish_attempt(
            attempt,
            terminal_status="invalid_response",
            provider_event_seen=True,
            terminal_seen=True,
            http_status=response.status_code,
            error_code=type(exc).__name__,
            committed=False,
        )
        await model_usage_audit.finish_logical_call(
            logical, terminal_status="failed", committed=False,
        )
        raise

    committed = response.status_code < 400 and bool(response_has_output(data))
    terminal_status = (
        "completed" if committed else "http_error" if response.status_code >= 400 else "invalid_response"
    )
    await model_usage_audit.finish_attempt(
        attempt,
        terminal_status=terminal_status,
        usage=model_usage_audit.provider_usage_from_response(data),
        response_id=model_usage_audit.provider_response_id(data),
        provider_event_seen=True,
        terminal_seen=True,
        http_status=response.status_code,
        committed=committed,
    )
    await model_usage_audit.finish_logical_call(
        logical,
        terminal_status="completed" if committed else "failed",
        selected_attempt_id=(attempt.attempt_id if attempt is not None and committed else ""),
        committed=committed,
    )
    return response, data


def _custom_ocr_has_output(data: dict) -> bool:
    return bool(data.get("text") or data.get("result") or data.get("content"))


def _chat_completion_has_output(data: dict) -> bool:
    return bool(((data.get("choices") or [{}])[0].get("message") or {}).get("content"))

MAX_TEXT_CHARS = settings.DOC_PARSE_MAX_TEXT_CHARS  # 附件正文截断上限（可经 DOC_PARSE_MAX_TEXT_CHARS 调）

# 多模态视觉识别默认提示词：不止 OCR 提字，而是完整描述图片内容，让文本型主对话模型
# 仅凭这段描述就能回答用户关于图片的各种问题（是什么、里面写了啥、数据、场景等）。
DEFAULT_VISION_PROMPT = (
    "请仔细观察这张图片并用中文完整描述其内容，覆盖："
    "①图中所有文字（原样转录，含标题、表格、标注）；"
    "②主要物体、人物、场景；③图表/表格中的数据与关系；④其他关键细节。"
    "力求准确详尽，使他人仅凭你的描述即可回答关于此图的问题；只输出描述本身，不要寒暄。"
)


async def _resolve_vision_runtime(conf: dict, newapi_key: str) -> Optional[dict]:
    """multimodal_model 策略下视觉调用用哪份 (base_url, api_key, model)。

    2026-09-19 用户拍板「OCR 不需要，我们用的是多模态模型」：扫描件 PDF / 文档内嵌图这些
    进知识库、文档解析的路径仍要把图变成文字，但**视觉模型就是平台对话模型**，不该再让
    管理员在「图片识别」里另填一套地址、密钥、模型名（该管理页 tab 已去掉）。

    优先级：
    1. ocr 里显式填齐的独立视觉端点（visionBaseUrl + visionApiKey + model 三者齐全）——
       管理员经 API 明确配置的仍优先；
    2. 平台对话模型连接（model_connection 平台级 runtime：管理员在「对话模型」里配的那份）；
    3. 旧路径兜底：当前请求绑定的网关 + 调用方带来的 key + ocr 里的模型名（三者都要有）。
    都拿不到返回 None，调用方给出明确的失败原因。
    """
    model = str(conf.get("model") or "").strip()
    vision_base = str(conf.get("visionBaseUrl") or "").strip().rstrip("/")
    vision_key = str(conf.get("visionApiKey") or "").strip()
    if model and vision_base and vision_key:
        return {"base_url": vision_base, "api_key": vision_key, "model": model, "source": "ocr_config"}
    try:
        from app.services.platform import model_connection
        platform = await model_connection.runtime_platform()
    except Exception:  # noqa: BLE001
        logger.info("读取平台对话模型连接失败，视觉识别退回旧网关路径", exc_info=True)
        platform = None
    if platform and platform.get("base_url") and platform.get("api_key") and platform.get("model"):
        return {
            "base_url": str(platform["base_url"]).rstrip("/"),
            "api_key": str(platform["api_key"]),
            "model": str(platform["model"]),
            "source": "platform_model",
        }
    if model and newapi_key:
        return {
            "base_url": get_model_base_url().rstrip("/"), "api_key": newapi_key,
            "model": model, "source": "gateway",
        }
    return None


_TEXT_EXTS = {
    "txt", "md", "markdown", "csv", "tsv", "json", "log", "yaml", "yml", "xml", "html", "htm",
    "py", "js", "ts", "tsx", "jsx", "java", "go", "rs", "c", "cpp", "h", "sh", "sql", "ini", "conf",
}
_IMAGE_EXTS = {"png", "jpg", "jpeg", "webp", "bmp", "gif"}
_IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "webp": "image/webp", "bmp": "image/bmp", "gif": "image/gif",
}
# 文档内嵌图片视觉识别：单文档最多识别几张（控上传时延与视觉调用成本）、
# 小于此字节的多半是图标/装饰（跳过）。仅在平台已启用 OCR（视觉）时才跑。
# PPT 天然图片型（满屏截图/一页一张配图很常见），单独放宽上限；docx/pdf 保持保守值。
MAX_EMBEDDED_IMAGES = 6
MAX_EMBEDDED_IMAGES_PPTX = 25
MIN_EMBEDDED_IMAGE_BYTES = 3000
_RASTER_EXTS = {"png", "jpg", "jpeg", "webp", "bmp", "gif"}

# docx/pptx/xlsx 本质是 zip：解析前先过解压体积闸（zip_guard），否则一个 <15MB 的恶意
# xlsx 能解出数 GB XML 把 worker 打爆（单层 DEFLATE 压缩比可达约 1000:1）。阈值取「正常
# Office 文档远够用、炸弹必被挡」的量级：真实百页 PPT 解压后普遍 < 200MB。
_ZIP_OFFICE_EXTS = {"docx", "pptx", "xlsx", "xlsm"}
DOC_ZIP_MAX_TOTAL_BYTES = 200 * 1024 * 1024
DOC_ZIP_MAX_ENTRY_BYTES = 100 * 1024 * 1024
DOC_ZIP_MAX_ENTRIES = 10_000


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_TEXT_CHARS:
        return text, False
    return text[:MAX_TEXT_CHARS] + f"\n...(已截断，原文共 {len(text)} 字符)", True


def _parse_pdf(content: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            continue
    return "\n".join(p for p in parts if p), len(reader.pages)


# 扫描版 PDF 判定：文字层平均每页不足该字符数视为无文字层（扫描件常见全空或仅剩页码）
SCAN_PDF_CHARS_PER_PAGE = 15
# 扫描版 PDF 最多 OCR 的页数：每页一次端点调用，控上传时延；超出部分明确标注
MAX_SCAN_PAGES = 15
_SCAN_RENDER_SCALE = 2.0  # 72dpi 基准 ×2 ≈ 144dpi，兼顾 OCR 精度与图片体积


def _render_pdf_pages(
    content: bytes, limit: int, *, scale: float | None = None,
) -> tuple[list[bytes], int]:
    """把 PDF 前 limit 页渲染为 PNG 字节序列。返回 (页面PNG列表, 总页数)。"""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content)
    try:
        total = len(pdf)
        pages: list[bytes] = []
        render_scale = float(scale or _SCAN_RENDER_SCALE)
        for i in range(min(total, limit)):
            buf = io.BytesIO()
            pdf[i].render(scale=render_scale).to_pil().save(buf, format="PNG")
            pages.append(buf.getvalue())
        return pages, total
    finally:
        pdf.close()


async def _ocr_pdf_scanned(
    content: bytes,
    *,
    newapi_key: str = "",
    audit_context: Optional[dict] = None,
) -> tuple[str, str, Optional[str]]:
    """扫描版 PDF（无文字层）：逐页渲染成 PNG 交 OCR 识别。

    优先走 `endpointUrl`（与图片的 strategy 解耦：图片可走视觉模型转述，文档 OCR
    仍用自建端点——2026-07-10 用户拍板的分工）。自建端点没填、但平台启用了
    multimodal_model 策略时，逐页交视觉模型转述（复用 _describe_image，与图片同一条
    链路）：线上只配了视觉模型、没有自建端点，此前扫描件在这里直接报「未配置 OCR
    端点」，管理页刚配好的视觉模型对 PDF 扫描件完全不起作用。两者都没有才返回明确提示。
    返回 (text, status, note)：status=ok/partial/failed——占位提示文本继续喂给模型
    （模型需要知道读不到的原因），status 供上层结构化下发前端（附件置信度）。
    """
    conf = await cfg.get_ocr_config()
    url = str(conf.get("endpointUrl") or "").strip()
    # 视觉模型名不再是必填：ocr 里没填时由 _resolve_vision_runtime 回落到平台对话模型。
    use_vision = (
        not url
        and bool(conf.get("enabled"))
        and conf.get("strategy") == "multimodal_model"
    )
    if not url and not use_vision:
        return "（扫描版 PDF 无文字层，且未配置 OCR 端点，无法提取文字）", "failed", "扫描版 PDF 未配置 OCR 端点"
    try:
        pages, total = await asyncio.to_thread(_render_pdf_pages, content, MAX_SCAN_PAGES)
    except Exception as e:  # noqa: BLE001
        logger.warning("扫描版 PDF 渲染失败: %s", e)
        return f"（扫描版 PDF 渲染失败：{str(e)[:120]}）", "failed", "扫描版 PDF 渲染失败"

    headers = {}
    if conf.get("apiKey"):
        headers["Authorization"] = f"Bearer {conf['apiKey']}"
    sem = asyncio.Semaphore(3)

    async def ocr_page(client: httpx.AsyncClient, idx: int, png: bytes) -> str:
        async with sem:
            if use_vision:
                # _describe_image 失败返回 ""，这里统一成与自建端点一致的失败占位
                described = await _describe_image(
                    png, "png", newapi_key, conf, audit_context=audit_context,
                )
                return described if described.strip() else f"（第 {idx + 1} 页 OCR 失败）"
            try:
                file_payload = {"file": (f"page{idx + 1}.png", png, "image/png")}
                _resp, data = await _audited_document_model_post(
                    client,
                    url,
                    model=str(conf.get("model") or "custom_ocr"),
                    purpose_detail="document_scan_ocr",
                    provider_api_key=str(conf.get("apiKey") or ""),
                    headers=headers,
                    wire_payload={"files": file_payload},
                    files=file_payload,
                    response_has_output=_custom_ocr_has_output,
                    audit_context=audit_context,
                )
                return str(data.get("text") or data.get("result") or data.get("content") or "")
            except Exception as e:  # noqa: BLE001
                logger.warning("扫描版 PDF 第 %s 页 OCR 失败: %s", idx + 1, e)
                return f"（第 {idx + 1} 页 OCR 失败）"

    async with httpx.AsyncClient(timeout=60) as client:
        texts = await asyncio.gather(*(ocr_page(client, i, p) for i, p in enumerate(pages)))
    body = "\n\n".join(f"【第 {i + 1} 页】\n{t.strip()}" for i, t in enumerate(texts))
    failed_pages = sum(1 for t in texts if t.startswith("（第 ") and t.endswith("页 OCR 失败）"))
    if total > len(pages):
        body += f"\n...(共 {total} 页，仅识别前 {len(pages)} 页)"
    text = f"【扫描版 PDF（OCR 识别）】\n{body}"
    if failed_pages >= len(texts) and texts:
        return text, "failed", "全部页面 OCR 失败"
    if failed_pages or total > len(pages):
        parts = []
        if failed_pages:
            parts.append(f"{failed_pages} 页 OCR 失败")
        if total > len(pages):
            parts.append(f"共 {total} 页仅识别前 {len(pages)} 页")
        return text, "partial", "、".join(parts)
    return text, "ok", None


def _parse_docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in document.paragraphs if p.text)


def _parse_pptx(content: bytes) -> str:
    from pptx import Presentation

    pages: list[str] = []
    for index, slide in enumerate(Presentation(io.BytesIO(content)).slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = str(shape.text_frame.text or "").strip()
                if value:
                    parts.append(value)
            if getattr(shape, "has_table", False):
                rows = ["\t".join(str(cell.text or "").strip() for cell in row.cells) for row in shape.table.rows]
                parts.extend(row for row in rows if row.strip())
        if parts:
            pages.append(f"【第 {index} 页】\n" + "\n".join(parts))
    return "\n\n".join(pages)


# Excel 解析护栏：附件正文预算有限（MAX_TEXT_CHARS），超大表只取每张工作表头部并如实标注；
# 需要全量统计/透视的场景由主对话引导走沙箱 execute_in_sandbox
MAX_XLSX_ROWS_PER_SHEET = 300
MAX_XLSX_COLS = 40


def _xlsx_cell(value: object) -> str:
    if value is None:
        return ""
    # 浮点整数值（Excel 数字单元格默认 float）显示成整数，避免满表 1.0/2.0
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())  # 压掉换行/连续空白，保持一格一列


def _parse_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    # data_only：公式单元格取缓存计算值（无缓存则为空）；read_only：流式读，防超大表撑爆内存
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: list[str] = []
    try:
        for ws in workbook.worksheets:
            rows: list[str] = []
            capped = False
            for row in ws.iter_rows(values_only=True):
                cells = [_xlsx_cell(v) for v in row[:MAX_XLSX_COLS]]
                if not any(cells):
                    continue
                rows.append(" | ".join(cells).rstrip())
                if len(rows) >= MAX_XLSX_ROWS_PER_SHEET:
                    capped = True
                    break
            if not rows:
                continue
            head = f"【工作表：{ws.title}】"
            if capped:
                head += f"（超长，仅展示前 {MAX_XLSX_ROWS_PER_SHEET} 个非空行）"
            sheets.append(head + "\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


async def _ocr_image(
    content: bytes,
    ext: str,
    newapi_key: str,
    *,
    audit_context: Optional[dict] = None,
) -> tuple[str, str, Optional[str]]:
    """图片 OCR/视觉识别。返回 (text, status, note)，占位提示照旧进 text 喂模型。"""
    conf = await cfg.get_ocr_config()
    if not conf.get("enabled") or conf.get("strategy") == "none":
        return "（图片已上传，但平台未启用 OCR，无法提取文字）", "failed", "平台未启用 OCR"

    strategy = conf.get("strategy")
    if strategy == "custom_endpoint":
        url = str(conf.get("endpointUrl") or "").strip()
        if not url:
            return "（OCR 端点未配置）", "failed", "OCR 端点未配置"
        headers = {}
        if conf.get("apiKey"):
            headers["Authorization"] = f"Bearer {conf['apiKey']}"
        try:
            async with httpx.AsyncClient(timeout=60) as client:
                file_payload = {
                    "file": (f"image.{ext}", content, _IMAGE_MIME.get(ext, "image/png"))
                }
                _resp, data = await _audited_document_model_post(
                    client,
                    url,
                    model=str(conf.get("model") or "custom_ocr"),
                    purpose_detail="document_image_ocr",
                    provider_api_key=str(conf.get("apiKey") or ""),
                    headers=headers,
                    wire_payload={"files": file_payload},
                    files=file_payload,
                    response_has_output=_custom_ocr_has_output,
                    audit_context=audit_context,
                )
            text = str(data.get("text") or data.get("result") or data.get("content") or "")[:MAX_TEXT_CHARS]
            if not text:
                return "（OCR 端点未返回文字）", "failed", "OCR 端点未返回文字"
            return text, "ok", None
        except Exception as e:  # noqa: BLE001
            logger.warning("自建 OCR 端点识别失败: %s", e)
            return f"（OCR 识别失败：{str(e)[:120]}）", "failed", "OCR 识别失败"

    if strategy == "multimodal_model":
        runtime = await _resolve_vision_runtime(conf, newapi_key)
        if runtime is None:
            return (
                "（视觉识别不可用：平台未配置对话模型，也没有独立的视觉模型端点）",
                "failed", "平台未配置对话模型",
            )
        base_url, api_key, model = runtime["base_url"], runtime["api_key"], runtime["model"]

        prompt = str(conf.get("visionPrompt") or "").strip() or DEFAULT_VISION_PROMPT
        mime = _IMAGE_MIME.get(ext, "image/png")
        data_url = f"data:{mime};base64,{base64.b64encode(content).decode()}"
        payload = {
            "model": model,
            "messages": [{
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": data_url}},
                ],
            }],
        }
        try:
            async with httpx.AsyncClient(timeout=90) as client:
                _resp, data = await _audited_document_model_post(
                    client,
                    f"{base_url}/chat/completions",
                    model=model,
                    purpose_detail="document_image_vision",
                    provider_api_key=api_key,
                    headers={"Authorization": f"Bearer {api_key}"},
                    wire_payload=payload,
                    json_payload=payload,
                    response_has_output=_chat_completion_has_output,
                    audit_context=audit_context,
                )
            text = str(((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "")[:MAX_TEXT_CHARS]
            # 加来源标注：让主对话模型知道这段是「视觉模型看图得来」，而非用户直接输入
            if text:
                return f"【图片内容（由视觉模型识别）】\n{text}", "ok", None
            return "（视觉模型未返回内容）", "failed", "视觉模型未返回内容"
        except Exception as e:  # noqa: BLE001
            logger.warning("多模态 OCR 识别失败: %s", e)
            return f"（图片识别失败：{str(e)[:120]}）", "failed", "图片识别失败"

    return f"（未知 OCR 策略：{strategy}）", "failed", "未知 OCR 策略"


def _img_ext_from_name(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else "png"


async def _describe_image(
    content: bytes,
    ext: str,
    newapi_key: str,
    conf: dict,
    *,
    audit_context: Optional[dict] = None,
) -> str:
    """给定已解析的 OCR 配置，返回图片识别文本（原始，无【】包裹）；失败/不支持返回 ""。
    供文档内嵌图片复用（不复用 _ocr_image，以免其面向直传图片的友好报错混进文档正文）。"""
    strategy = conf.get("strategy")
    if strategy == "custom_endpoint":
        url = str(conf.get("endpointUrl") or "").strip()
        if not url:
            return ""
        headers = {}
        if conf.get("apiKey"):
            headers["Authorization"] = f"Bearer {conf['apiKey']}"
        try:
            async with httpx.AsyncClient(timeout=60) as client:
                file_payload = {
                    "file": (f"image.{ext}", content, _IMAGE_MIME.get(ext, "image/png"))
                }
                _resp, data = await _audited_document_model_post(
                    client,
                    url,
                    model=str(conf.get("model") or "custom_ocr"),
                    purpose_detail="document_embedded_image_ocr",
                    provider_api_key=str(conf.get("apiKey") or ""),
                    headers=headers,
                    wire_payload={"files": file_payload},
                    files=file_payload,
                    response_has_output=_custom_ocr_has_output,
                    audit_context=audit_context,
                )
            return str(data.get("text") or data.get("result") or data.get("content") or "")[:MAX_TEXT_CHARS]
        except Exception as e:  # noqa: BLE001
            logger.warning("内嵌图片自建 OCR 失败: %s", e)
            return ""
    if strategy == "multimodal_model":
        runtime = await _resolve_vision_runtime(conf, newapi_key)
        if runtime is None:
            return ""
        base_url, api_key, model = runtime["base_url"], runtime["api_key"], runtime["model"]
        prompt = str(conf.get("visionPrompt") or "").strip() or DEFAULT_VISION_PROMPT
        mime = _IMAGE_MIME.get(ext, "image/png")
        data_url = f"data:{mime};base64,{base64.b64encode(content).decode()}"
        payload = {"model": model, "messages": [{
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": data_url}},
            ],
        }]}
        try:
            async with httpx.AsyncClient(timeout=90) as client:
                _resp, data = await _audited_document_model_post(
                    client,
                    f"{base_url}/chat/completions",
                    model=model,
                    purpose_detail="document_embedded_image_vision",
                    provider_api_key=api_key,
                    headers={"Authorization": f"Bearer {api_key}"},
                    wire_payload=payload,
                    json_payload=payload,
                    response_has_output=_chat_completion_has_output,
                    audit_context=audit_context,
                )
            return str(((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "")[:MAX_TEXT_CHARS]
        except Exception as e:  # noqa: BLE001
            logger.warning("内嵌图片多模态识别失败: %s", e)
            return ""
    return ""


# 内嵌位图抽取预算：配图是「锦上添花」的可选增强，单独给一份更紧的实读预算——
# 中央目录声明值可伪造，这里按真实解出的字节计数，越线即停（已取到的照常识别）。
_MEDIA_BUDGET_TOTAL_BYTES = 64 * 1024 * 1024
_MEDIA_BUDGET_ENTRY_BYTES = 32 * 1024 * 1024


def _extract_zip_media(content: bytes, prefix: str, label: str) -> list[tuple[str, bytes]]:
    """从 OOXML（docx/pptx 本质是 zip）的 media 目录取内嵌位图。返回 [(文件名, 字节)]。"""
    out: list[tuple[str, bytes]] = []
    budget = zip_guard.ZipReadBudget(
        label=label,
        max_total_bytes=_MEDIA_BUDGET_TOTAL_BYTES,
        max_entry_bytes=_MEDIA_BUDGET_ENTRY_BYTES,
    )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            for name in zf.namelist():
                low = name.lower()
                if low.startswith(prefix) and _img_ext_from_name(low) in _RASTER_EXTS:
                    out.append((name.rsplit("/", 1)[-1], budget.read(zf, name)))
    except zip_guard.ZipBombError as e:
        # 配图超预算不致命：正文文字已单独解析，这里保留已取到的图，不再往下读
        logger.warning("提取%s内嵌图片中止（体积超限）: %s", label, e.reason)
    except Exception as e:  # noqa: BLE001
        logger.warning("提取%s内嵌图片失败: %s", label, e)
    return out


def _extract_docx_images(content: bytes) -> list[tuple[str, bytes]]:
    """从 docx（本质是 zip）的 word/media/ 取内嵌位图字节。返回 [(文件名, 字节)]。"""
    return _extract_zip_media(content, "word/media/", "docx")


def _extract_pptx_images(content: bytes) -> list[tuple[str, bytes]]:
    """从 pptx 的 ppt/media/ 取出内嵌位图；与文本抽取共用同一视觉识别管线。"""
    return _extract_zip_media(content, "ppt/media/", "pptx")


def _extract_pdf_images(content: bytes, limit: int) -> list[tuple[str, bytes]]:
    """从 PDF 各页取内嵌位图（pypdf page.images），累计到 limit 张即止。"""
    from pypdf import PdfReader

    out: list[tuple[str, bytes]] = []
    try:
        reader = PdfReader(io.BytesIO(content))
        for page in reader.pages:
            try:
                for img in page.images:
                    ext = _img_ext_from_name(img.name or "img.png")
                    if ext in _RASTER_EXTS and img.data:
                        out.append((img.name or f"img{len(out)}.{ext}", img.data))
                        if len(out) >= limit:
                            return out
            except Exception:  # noqa: BLE001
                continue  # 个别页图片解码失败不影响其余
    except Exception as e:  # noqa: BLE001
        logger.warning("提取 pdf 内嵌图片失败: %s", e)
    return out


async def _ocr_embedded_images(
    images: list[tuple[str, bytes]], newapi_key: str, doc_label: str,
    max_images: int = MAX_EMBEDDED_IMAGES,
    *,
    audit_context: Optional[dict] = None,
) -> tuple[str, str, Optional[str]]:
    """对文档内嵌图片逐张视觉识别，返回 (text, status, note)。

    去重（重复 logo 只识别一次）+ 跳过过小装饰图 + 数量上限 + 并发限制（同扫描 PDF）。
    """
    if not images:
        return "", "ok", None

    seen: set[str] = set()
    picked: list[tuple[str, bytes]] = []
    omitted_meaningful = 0
    for name, blob in images:
        if len(blob) < MIN_EMBEDDED_IMAGE_BYTES:
            continue
        digest = hashlib.sha1(blob).hexdigest()
        if digest in seen:
            continue
        seen.add(digest)
        if len(picked) < max_images:
            picked.append((name, blob))
        else:
            omitted_meaningful += 1
    if not picked:
        return "", "ok", None

    conf = await cfg.get_ocr_config()
    if not conf.get("enabled") or conf.get("strategy") == "none":
        return "", "partial", f"{doc_label}包含图片，但平台未启用视觉识别"

    sem = asyncio.Semaphore(3)

    async def describe(name: str, blob: bytes) -> str:
        async with sem:
            return (
                await _describe_image(
                    blob,
                    _img_ext_from_name(name),
                    newapi_key,
                    conf,
                    audit_context=audit_context,
                )
            ).strip()

    results = await asyncio.gather(*(describe(n, b) for n, b in picked))
    blocks = [f"【{doc_label}内嵌图片 {i + 1}】\n{d}" for i, d in enumerate(results) if d]
    if not blocks:
        return "", "partial", f"{doc_label}内嵌图片识别失败"
    body = f"【{doc_label}中的图片内容（由视觉模型识别）】\n\n" + "\n\n".join(blocks)
    omitted = omitted_meaningful > 0
    failed = sum(1 for result in results if not result)
    if omitted:
        body += f"\n\n（{doc_label}配图较多，仅对其中前 {len(picked)} 张做了视觉识别；正文文字已全部读取）"
    if failed or omitted:
        notes = []
        if failed:
            notes.append(f"{failed} 张配图识别失败")
        if omitted:
            # 「张配图」而非「张」——避免被误读成「只读了前 N 页幻灯片」（正文文字其实已全读）
            notes.append(f"仅识别前 {len(picked)} 张配图")
        return body, "partial", "、".join(notes)
    return body, "ok", None


def _merge_read_status(
    current: str, current_note: Optional[str], incoming: str, incoming_note: Optional[str]
) -> tuple[str, Optional[str]]:
    if incoming == "failed" or (incoming == "partial" and current == "ok"):
        return incoming, incoming_note
    return current, current_note


_VISUAL_SECTION_MARK = "中的图片内容（由视觉模型识别）】"


def _truncate_document(text: str) -> tuple[str, bool]:
    """长文档截断时保留末尾的内嵌图片描述，避免图片已识别却在入模前被切掉。"""
    if len(text) <= MAX_TEXT_CHARS:
        return text, False
    marker = text.find(_VISUAL_SECTION_MARK)
    if marker < 0:
        return _truncate(text)
    marker = text.rfind("【", 0, marker + 1)
    visual = text[marker:]
    visual_budget = min(8000, len(visual))
    head_budget = MAX_TEXT_CHARS - visual_budget - 80
    clipped = (
        text[:head_budget].rstrip()
        + f"\n...(文档文字已截断，原文共 {len(text)} 字符)\n\n"
        + visual[:visual_budget]
    )
    return clipped, True


async def parse_upload(
    filename: str,
    content: bytes,
    *,
    newapi_key: str = "",
    ocr_embedded_images: bool = True,
    ocr_visual: bool = True,
    audit_context: Optional[dict] = None,
) -> dict:
    """解析上传文件为文本。返回 {filename, kind, text, chars, truncated}。

    ocr_embedded_images：docx/pdf/pptx 内嵌图片是否走视觉识别补进正文（默认开）。对话上下文
    路径（/chat/upload、我的文件带入对话、read_file）用默认开；文件「预览」为保响应快，
    调用方显式传 False 关闭（预览只看文字，不为几张图等十几秒视觉调用）。
    ocr_visual：独立图片 / 扫描 PDF 是否走 OCR 视觉模型。主对话当前模型已是多模态时
    应关掉——像素会在发送时以 image_url 直传，上传阶段不必再等一轮视觉调用。

    返回附带结构化读取置信度 status: ok（完整）/ partial（截断或部分页失败）/ failed
    （未提取出可用内容）+ note（简短原因）。此前失败只以占位文本混在 text 里（HTTP 200），
    前端与回答降级逻辑都无从感知——结构化字段是附件可信度链路的事实源。"""
    ext = _ext(filename)
    kind = "text"
    text = ""
    status = "ok"
    note: Optional[str] = None

    # 旧版 Office 二进制（OLE 复合文档：.ppt/.doc/.xls，97-2003）——python-pptx/python-docx/
    # openpyxl 只能开新版 OOXML（.pptx/.docx/.xlsx 本质是 zip），agent-api 又没有 LibreOffice
    # 做转换（只有沙箱里有）。此前一律抛到下面 except 报模糊的「文件解析失败」，用户无从下手。
    # 改为给可执行提示：另存为新版格式再上传（新版解析已支持，实测可读）。（用户报告 2026-07-17）
    _legacy_modern = {"ppt": "pptx", "doc": "docx", "xls": "xlsx"}.get(ext)
    _is_ole_binary = content[:8].startswith(b"\xd0\xcf\x11\xe0")  # OLE 魔数（含改名成 .pptx 的旧文件）
    if _legacy_modern or (_is_ole_binary and ext in ("pptx", "docx", "xlsx", "xlsm", "ppt", "doc", "xls", "")):
        target = _legacy_modern or (
            "pptx" if "ppt" in ext else "docx" if "doc" in ext else "xlsx" if "xls" in ext else "pptx / docx / xlsx"
        )
        return {
            "filename": filename, "kind": "text",
            "text": (f"（《{filename}》是旧版 Office 二进制格式（97-2003），上传解析暂不支持；"
                     f"请在 Office / WPS 里「另存为 .{target}」后重新上传，或导出为 PDF。）"),
            "chars": 0, "truncated": False, "status": "failed",
            "note": f"旧版二进制格式，请另存为 .{target} 或 PDF 再上传",
        }

    # 解压炸弹闸：docx/pptx/xlsx 解析前先按 zip 中央目录声明的解压后体积判上限，
    # 越线直接 failed（别等 python-docx/openpyxl 真去解，那时内存已经没了）。
    if ext in _ZIP_OFFICE_EXTS:
        try:
            zip_guard.ensure_zip_bytes_within_limits(
                content,
                label=f"《{filename}》",
                max_total_bytes=DOC_ZIP_MAX_TOTAL_BYTES,
                max_entry_bytes=DOC_ZIP_MAX_ENTRY_BYTES,
                max_entries=DOC_ZIP_MAX_ENTRIES,
            )
        except zip_guard.ZipBombError as e:
            logger.warning("拒绝解析疑似解压炸弹文件 %s: %s", filename, e.reason)
            return {
                "filename": filename, "kind": "text", "text": f"（{e.message}）",
                "chars": 0, "truncated": False, "status": "failed", "note": e.reason,
            }

    # 下面的 docx/pptx/xlsx/pdf 解析全是**同步 CPU 密集**（zip 解压 + XML 解析 + 图片解码）：
    # 必须 asyncio.to_thread 丢出事件循环，否则单个几十 MB 的文档就能把整个 worker 的事件
    # 循环独占几十秒——同进程所有 SSE 流一起卡住（P0）。这些函数都只吃 bytes 入参、内部
    # 局部 import、无跨线程共享可变状态，换线程执行是安全的。
    try:
        if ext in _TEXT_EXTS or not ext:
            text = content.decode("utf-8", errors="ignore")
        elif ext == "pdf":
            kind = "pdf"
            text, npages = await asyncio.to_thread(_parse_pdf, content)
            # 文字层过薄 → 按扫描件处理（整页渲染 OCR，已含图）；OCR 结果更充实才替换，
            # 避免端点故障时反把已有的少量真实文字冲掉
            if len(text.strip()) < SCAN_PDF_CHARS_PER_PAGE * max(1, npages):
                if ocr_visual:
                    ocr_text, ocr_status, ocr_note = await _ocr_pdf_scanned(
                        content, newapi_key=newapi_key, audit_context=audit_context,
                    )
                    if len(ocr_text.strip()) > len(text.strip()):
                        text = ocr_text
                        status, note = ocr_status, ocr_note
                    elif not text.strip():
                        status, note = "failed", ocr_note or "PDF 无文字层且 OCR 不可用"
                elif not text.strip():
                    text = f"（《{filename}》为扫描件/页图，已交给多模态模型直接查看，未做 OCR）"
                    status, note = "ok", "扫描件将由多模态模型直接查看页图"
            elif ocr_embedded_images:
                # 有文字层但可能夹着照片/图表：抽出内嵌位图逐张视觉识别，补进正文
                img_text, img_status, img_note = await _ocr_embedded_images(
                    await asyncio.to_thread(_extract_pdf_images, content, MAX_EMBEDDED_IMAGES + 1),
                    newapi_key, "PDF", audit_context=audit_context,
                )
                if img_text:
                    text = f"{text}\n\n{img_text}" if text.strip() else img_text
                status, note = _merge_read_status(status, note, img_status, img_note)
        elif ext == "docx":
            kind = "docx"
            text = await asyncio.to_thread(_parse_docx, content)
            # Word 里塞的图片（照片/截图/图表）走视觉识别补进正文——否则只看得到文字
            if ocr_embedded_images:
                img_text, img_status, img_note = await _ocr_embedded_images(
                    await asyncio.to_thread(_extract_docx_images, content),
                    newapi_key,
                    "文档",
                    audit_context=audit_context,
                )
                if img_text:
                    text = f"{text}\n\n{img_text}" if text.strip() else img_text
                status, note = _merge_read_status(status, note, img_status, img_note)
        elif ext == "pptx":
            kind = "pptx"
            text = await asyncio.to_thread(_parse_pptx, content)
            if ocr_embedded_images:
                img_text, img_status, img_note = await _ocr_embedded_images(
                    await asyncio.to_thread(_extract_pptx_images, content), newapi_key, "演示文稿",
                    max_images=MAX_EMBEDDED_IMAGES_PPTX,
                    audit_context=audit_context,
                )
                if img_text:
                    text = f"{text}\n\n{img_text}" if text.strip() else img_text
                status, note = _merge_read_status(status, note, img_status, img_note)
        elif ext in _IMAGE_EXTS:
            kind = "image"
            if ocr_visual:
                text, status, note = await _ocr_image(
                    content, ext, newapi_key, audit_context=audit_context,
                )
            else:
                text, status, note = "", "ok", "原图已保存，将由多模态模型直接查看"
        elif ext in ("xlsx", "xlsm"):
            # .xls 旧二进制在上方 OLE 分支已给「另存为 .xlsx」引导，这里只处理新版 OOXML
            kind = "xlsx"
            text = await asyncio.to_thread(_parse_xlsx, content)
        else:
            # 兜底尝试按文本解码
            text = content.decode("utf-8", errors="ignore")
            if not text:
                text = f"（不支持的文件类型：.{ext}）"
                status, note = "failed", f"不支持的文件类型 .{ext}"
    except Exception as e:  # noqa: BLE001
        logger.warning("解析文件 %s 失败: %s", filename, e)
        text = f"（文件解析失败：{str(e)[:120]}）"
        status, note = "failed", "文件解析失败"

    # PG text 列不接受 NUL(0x00)：pdf 文字层/OCR 返回偶发夹带，入库前统一剔除，
    # 否则会话附件持久化失败，同会话后续轮次丢附件上下文。
    text, truncated = _truncate_document(text.replace("\x00", "").strip())
    if truncated and status == "ok":
        status, note = "partial", "内容超长已截断"
    # 文档类解析出空文本＝没读到任何内容（纯图片/空白 docx、无文字层且 OCR 未补上的 pdf），
    # 不得标 ok——否则模型和用户都以为「完整读取」了一个空字符串（审查 P1）
    if status == "ok" and not text and kind in ("docx", "pdf", "pptx", "xlsx") and ocr_visual:
        text = f"（《{filename}》未解析出文本内容，可能是空白文档或纯图片文档）"
        status, note = "failed", "未解析出文本内容（可能是空白或纯图片文档）"
    # partial 也可能空文本（纯图片文档且 OCR 未启用/全部识别失败）：同样不能带着空字符串
    # 下发——前端只拦 failed，partial+空文本会渲染成正常卡片、发送时又被空文本过滤静默丢弃
    if status == "partial" and not text:
        text = f"（《{filename}》未解析出文本内容：{note or '内嵌图片识别失败'}）"
        status = "failed"
        note = note or "未解析出文本内容"
    return {
        "filename": filename, "kind": kind, "text": text, "chars": len(text),
        "truncated": truncated, "status": status, "note": note,
    }
