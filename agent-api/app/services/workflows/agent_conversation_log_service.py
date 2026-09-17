"""Read model and XLSX output for workflow-agent conversation logs."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time, timedelta
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import or_, select

from app.models import ChatMessage, ChatThread
from app.services.platform.user_display_name import load_user_display_names
from app.services.workflows.agent_metrics_service import METRIC_RANGE_KEYS, resolve_metrics_window


SHANGHAI_TZ = ZoneInfo("Asia/Shanghai")
LOG_STATUSES = {"completed", "partial", "failed", "cancelled", "interrupted", "unknown"}
EXPORT_HEADERS = [
    "应用名称", "会话标题", "用户姓名", "会话状态", "消息数", "最近消息时间", "点赞数", "点踩数",
    "提问时间", "用户问题", "回答时间", "智能体回答", "回答状态", "用户反馈",
]
MAX_EXPORT_QUESTION_ANSWER_ROWS = 100_000


@dataclass(frozen=True)
class ConversationLogFilters:
    start_at: datetime | None
    end_at: datetime | None
    range_key: str
    status: str | None
    keyword: str
    page_no: int
    page_size: int


def _parse_datetime(value: str | None, name: str) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"{name} 必须是 ISO 时间") from exc
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(SHANGHAI_TZ).replace(tzinfo=None)
    return parsed


def normalize_conversation_log_filters(
    start_at: str | None = None,
    end_at: str | None = None,
    range_key: str = "last_7_days",
    status: str | None = None,
    keyword: str | None = None,
    page_no: int = 1,
    page_size: int = 20,
) -> ConversationLogFilters:
    now = datetime.now(SHANGHAI_TZ).replace(tzinfo=None)
    normalized_range = (range_key or "last_7_days").strip()
    if start_at or end_at:
        end = _parse_datetime(end_at, "endAt") or now
        start = _parse_datetime(start_at, "startAt") or end - timedelta(days=7)
        normalized_range = "custom"
    else:
        if normalized_range not in METRIC_RANGE_KEYS:
            raise ValueError("不支持的 range 参数")
        if normalized_range == "all_time":
            start = end = None
        else:
            window = resolve_metrics_window(normalized_range)
            start = datetime.combine(window.start, time.min)
            end = datetime.combine(window.end + timedelta(days=1), time.min)
    normalized_status = (status or "").strip() or None
    normalized_keyword = (keyword or "").strip()
    if normalized_status and normalized_status not in LOG_STATUSES:
        raise ValueError("不支持的状态筛选")
    if start is not None and end is not None and start >= end:
        raise ValueError("startAt 必须早于 endAt")
    if len(normalized_keyword) > 100:
        raise ValueError("关键词不能超过 100 个字符")
    if page_no < 1 or page_size < 1 or page_size > 100:
        raise ValueError("分页参数无效")
    return ConversationLogFilters(start, end, normalized_range, normalized_status, normalized_keyword, page_no, page_size)


def _status(value: str | None) -> str:
    return value if value in LOG_STATUSES - {"unknown"} else "unknown"


def _human_message(row: ChatMessage) -> bool:
    return row.role == "user" and row.sender_type in (None, "human")


def format_conversation_log_time(value: datetime | str | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return text.replace("T", " ")[:19]


async def query_conversation_logs(
    session: Any,
    app_id: str,
    filters: ConversationLogFilters,
    *,
    include_all: bool = False,
    include_transcript: bool = False,
) -> dict[str, Any]:
    assistant = ChatMessage
    statement = (
        select(assistant, ChatThread)
        .join(ChatThread, ChatThread.id == assistant.thread_id)
        .where(
            ChatThread.app_id == app_id,
            assistant.role == "assistant",
        )
        .order_by(assistant.created_at.desc(), assistant.id.desc())
    )
    if filters.start_at is not None and filters.end_at is not None:
        statement = statement.where(assistant.created_at >= filters.start_at, assistant.created_at < filters.end_at)
    if filters.status == "unknown":
        statement = statement.where(or_(assistant.status.is_(None), assistant.status == "unknown"))
    elif filters.status:
        statement = statement.where(assistant.status == filters.status)
    anchors = (await session.execute(statement)).all()
    if not anchors:
        return {"records": [], "total": 0}

    thread_ids = sorted({str(thread.id) for _, thread in anchors})
    rows = (await session.execute(
        select(ChatMessage)
        .where(ChatMessage.thread_id.in_(thread_ids))
        .order_by(ChatMessage.thread_id.asc(), ChatMessage.created_at.asc(), ChatMessage.id.asc())
    )).scalars().all()
    by_thread: dict[str, list[ChatMessage]] = {}
    for row in rows:
        by_thread.setdefault(str(row.thread_id), []).append(row)

    display_names = await load_user_display_names(session, [str(thread.user_id) for _, thread in anchors])
    records: list[dict[str, Any]] = []
    seen_thread_ids: set[str] = set()
    for anchor, thread in anchors:
        thread_id = str(thread.id)
        if thread_id in seen_thread_ids:
            continue
        seen_thread_ids.add(thread_id)
        thread_rows = by_thread.get(thread_id, [])
        upvotes = sum(1 for row in thread_rows if row.feedback == "up")
        downvotes = sum(1 for row in thread_rows if row.feedback == "down")
        record = {
            "id": thread_id,
            "threadId": thread_id,
            "title": thread.title or "新对话",
            "userId": str(thread.user_id),
            "username": display_names.get(str(thread.user_id), str(thread.user_id)),
            "status": _status(anchor.status),
            "messageCount": len(thread_rows),
            "lastMessageAt": format_conversation_log_time(anchor.created_at),
            "upvotes": upvotes,
            "downvotes": downvotes,
        }
        searchable = "\n".join([record["title"], record["username"], *(row.content or "" for row in thread_rows)]).lower()
        if filters.keyword and filters.keyword.lower() not in searchable:
            continue
        if include_transcript:
            record["messages"] = [
                {
                    "isUser": _human_message(row),
                    "role": row.role,
                    "content": row.content or "",
                    "status": _status(row.status) if row.role == "assistant" else None,
                    "feedback": row.feedback if row.feedback in {"up", "down"} else None,
                    "createdAt": format_conversation_log_time(row.created_at),
                }
                for row in thread_rows
            ]
        records.append(record)
    total = len(records)
    if not include_all:
        start_index = (filters.page_no - 1) * filters.page_size
        records = records[start_index:start_index + filters.page_size]
    return {"records": records, "total": total}


async def query_conversation_log_detail(session: Any, app_id: str, thread_id: str) -> dict[str, Any] | None:
    thread = (await session.execute(
        select(ChatThread).where(ChatThread.id == thread_id, ChatThread.app_id == app_id)
    )).scalar_one_or_none()
    if thread is None:
        return None
    rows = (await session.execute(
        select(ChatMessage)
        .where(ChatMessage.thread_id == thread_id)
        .order_by(ChatMessage.created_at.asc(), ChatMessage.id.asc())
    )).scalars().all()
    display_names = await load_user_display_names(session, [str(thread.user_id)])
    return {
        "threadId": str(thread.id),
        "title": thread.title or "新对话",
        "userId": str(thread.user_id),
        "username": display_names.get(str(thread.user_id), str(thread.user_id)),
        "messages": [
            {
                "id": row.id,
                "role": row.role,
                "content": row.content or "",
                "status": _status(row.status) if row.role == "assistant" else None,
                "feedback": row.feedback if row.feedback in {"up", "down"} else None,
                "createdAt": format_conversation_log_time(row.created_at),
            }
            for row in rows
        ],
    }


def excel_text(value: object) -> str:
    text = str(value or "")
    return f"'{text}" if text[:1] in {"=", "+", "-", "@"} else text


def _conversation_question_answer_rows(messages: list[dict[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    pending_question: dict[str, str] | None = None
    for message in messages:
        is_user = bool(message.get("isUser")) or message.get("role") == "user"
        if is_user:
            if pending_question is not None:
                rows.append(pending_question)
            pending_question = {
                "questionAt": str(message.get("createdAt") or ""),
                "question": str(message.get("content") or ""),
                "answerAt": "",
                "answer": "",
                "answerStatus": "",
                "feedback": "",
            }
            continue

        answer = {
            "answerAt": str(message.get("createdAt") or ""),
            "answer": str(message.get("content") or ""),
            "answerStatus": str(message.get("status") or ""),
            "feedback": "点赞" if message.get("feedback") == "up" else "点踩" if message.get("feedback") == "down" else "",
        }
        if pending_question is not None:
            pending_question.update(answer)
            rows.append(pending_question)
            pending_question = None
        else:
            rows.append({"questionAt": "", "question": "", **answer})
    if pending_question is not None:
        rows.append(pending_question)
    return rows or [{"questionAt": "", "question": "", "answerAt": "", "answer": "", "answerStatus": "", "feedback": ""}]


def build_conversation_log_xlsx(app_name: str, records: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "对话日志"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    question_answer_count = 0
    for record in records:
        turns = _conversation_question_answer_rows(record.get("messages") or [])
        question_answer_count += len(turns)
        if question_answer_count > MAX_EXPORT_QUESTION_ANSWER_ROWS:
            raise ValueError(f"导出问答不能超过 {MAX_EXPORT_QUESTION_ANSWER_ROWS} 条，请缩小筛选范围")
        start_row = sheet.max_row + 1
        session_values = [
            app_name, record.get("title"), record.get("username"), record.get("status"), record.get("messageCount"),
            format_conversation_log_time(record.get("lastMessageAt")), record.get("upvotes"), record.get("downvotes"),
        ]
        for index, turn in enumerate(turns):
            sheet.append([
                *(excel_text(value) if index == 0 else "" for value in session_values),
                excel_text(format_conversation_log_time(turn["questionAt"])), excel_text(turn["question"]), excel_text(format_conversation_log_time(turn["answerAt"])),
                excel_text(turn["answer"]), excel_text(turn["answerStatus"]), excel_text(turn["feedback"]),
            ])
        end_row = sheet.max_row
        if end_row > start_row:
            for column in range(1, 9):
                sheet.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)
        for row in sheet.iter_rows(min_row=start_row, max_row=end_row):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGHIJKLMN", (22, 24, 16, 12, 10, 22, 10, 10, 22, 42, 22, 52, 12, 12)):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
